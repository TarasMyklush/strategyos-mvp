"""Real PostgreSQL proof in a disposable, socket-only cluster; never business DBs."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timezone
from decimal import Decimal
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
from uuid import uuid4

from fastapi import FastAPI
from fastapi.testclient import TestClient
import psycopg
from psycopg import sql
import pytest

from strategyos_mvp import auth, dimensional_intent_sources as sources
from strategyos_mvp import dimensional_intent_store as store
from strategyos_mvp.dimensional_intent_api import router
from strategyos_mvp.dimensional_plan import Actuals, Plan, fingerprint
from strategyos_mvp.plan_decomposition import (DecompositionRequest, HistoricalDecompositionRequest,
                                               ObjectiveDecompositionRequest)
from strategyos_mvp.advisor_config import AdvisorConfiguration
from strategyos_mvp import advisor_config_store
from strategyos_mvp.tenant_structure import TenantStructureConfiguration
from strategyos_mvp import tenant_structure_store

FIXTURE = Path(__file__).parent / 'fixtures' / 'dimensional_plan'
TODAY = datetime.now(timezone.utc).date()


@pytest.fixture(scope='module')
def postgres():
    initdb, pg_ctl = shutil.which('initdb'), shutil.which('pg_ctl')
    if not initdb:
        candidates = sorted(Path('/usr/lib/postgresql').glob('*/bin/initdb'))
        if candidates:
            initdb = str(candidates[-1])
            pg_ctl = str(candidates[-1].with_name('pg_ctl'))
    if not initdb or not pg_ctl:
        pytest.skip('Local PostgreSQL initdb/pg_ctl required for isolated dimensional proof.')
    with tempfile.TemporaryDirectory(prefix='ki-') as directory:
        root = Path(directory)
        data, socket, log = root / 'db', root / 's', root / 'postgres.log'
        socket.mkdir()
        subprocess.run([initdb, '-D', str(data), '-U', 'intent_proof', '-A', 'trust', '--no-locale', '-E', 'UTF8'],
                       check=True, capture_output=True, timeout=30)
        # No TCP listener; no environment credentials or existing service are used.
        subprocess.run([pg_ctl, '-D', str(data), '-l', str(log), '-o',
                        f"-F -k {socket} -p 55439 -h ''", '-w', 'start'],
                       check=True, capture_output=True, timeout=30)
        try:
            yield {'host': str(socket), 'port': 55439, 'user': 'intent_proof'}
        finally:
            subprocess.run([pg_ctl, '-D', str(data), '-m', 'immediate', '-w', 'stop'],
                           check=True, capture_output=True, timeout=30)


@pytest.fixture
def setup(postgres, monkeypatch, tmp_path):
    database = 'intent_' + uuid4().hex
    with psycopg.connect(**postgres, dbname='postgres', autocommit=True) as conn:
        conn.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(database)))
    connect = lambda: psycopg.connect(**postgres, dbname=database)
    monkeypatch.setattr(store.state_store, 'database_connection', lambda: (connect(), None))
    config = replace(store.CONFIG, tenant_slug='intent-tenant', output_root=tmp_path / 'outputs')
    monkeypatch.setattr(store, 'CONFIG', config)
    monkeypatch.setattr(sources, 'CONFIG', config)
    store.initialize()
    with connect() as conn:
        conn.execute("CREATE TABLE strategyos_tenants(id integer primary key,slug text)")
        conn.execute("CREATE TABLE strategyos_source_systems(id integer primary key,tenant_id integer,source_key text)")
        conn.execute("CREATE TABLE strategyos_source_access_policies(source_system_id integer,allowed_roles text[],allowed_purposes text[],allowed_business_units text[],storage_allowed boolean,export_allowed boolean,effective_to timestamptz)")
        conn.execute("INSERT INTO strategyos_tenants VALUES(1,'intent-tenant')")
        conn.execute("INSERT INTO strategyos_source_systems VALUES(1,1,'intent-proof')")
        conn.execute("INSERT INTO strategyos_source_access_policies VALUES(1,%s,ARRAY['analysis','export'],ARRAY[]::text[],true,true,NULL)", (list(store.READ_ROLES),))
    pack = 'owned-pack'
    root = config.output_root / 'source_packs' / pack
    (root / 'raw').mkdir(parents=True)
    shutil.copy(FIXTURE / 'evidence.csv', root / 'raw')
    p, a = (json.loads((FIXTURE / f'{name}.json').read_text()) for name in ['plan', 'actuals'])
    p['company_id'] = a['company_id'] = config.tenant_slug
    period = {'start': f'{TODAY.year-1}-01-01', 'end': f'{TODAY.year-1}-01-31'}
    p['period'], a['period'] = deepcopy(period), deepcopy(period)
    p['effective_from'], p['effective_to'] = period['start'], f'{TODAY.year-1}-12-31'
    a['recorded_on'] = period['end']
    manifest = {'source_contract': {'source_key': 'intent-proof'}, 'source_pack_id': pack, 'tenant_context': {'tenant_id': config.tenant_slug},
                'manifest': [{'relative_path': 'evidence.csv', 'supported': True,
                              'sha256': p['cells'][0]['source']['sha256'], 'source_disposition': 'current_evidence'}]}
    (root / 'summary.json').write_text(json.dumps(manifest))
    principal = lambda role, subject: {'tenant_id': config.tenant_slug, 'role': role, 'subject': subject}
    admin = principal('tenant_admin', 'admin')
    operator = principal('operator', 'operator')
    executive = principal('executive', 'executive')
    fixture_structure = TenantStructureConfiguration.model_validate(fixture_structure_body())
    structure_record = tenant_structure_store.create(operator, fixture_structure)
    tenant_structure_store.approve(
        admin, fixture_structure.config_id, fixture_structure.version, structure_record['digest'],
        'Independently reviewed fixture organization and dimension structure.')
    p['business_unit'] = 'group'
    p['structure'] = {'config_id': fixture_structure.config_id, 'version': fixture_structure.version,
                      'digest': structure_record['digest']}
    app = FastAPI()
    app.include_router(router)
    current = {'principal': operator}
    app.dependency_overrides[auth.authenticate_request] = lambda: current['principal']
    with TestClient(app) as client:
        yield dict(p=p, a=a, pack=pack, root=root, manifest=manifest, admin=admin, operator=operator,
                   executive=executive, connect=connect, client=client, current=current, app=app)
    with psycopg.connect(**postgres, dbname='postgres', autocommit=True) as conn:
        conn.execute(sql.SQL('DROP DATABASE {}').format(sql.Identifier(database)))


def import_pair(s):
    plan = store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    store.import_actuals(s['operator'], Actuals.model_validate(s['a']), s['pack'])
    return plan


def approve(s, plan):
    store.set_ratifier(s['admin'], s['p']['plan_id'], s['executive']['subject'], True, 0)
    return store.ratify(s['executive'], s['p']['plan_id'], s['p']['version'], plan['digest'],
                        'Reviewed the targets, owners and source evidence.')


def analyse(s):
    return store.create_analysis(s['executive'], s['p']['plan_id'], s['p']['version'], s['a']['revision'], TODAY)


def structure_body(company='Example Holdings', config_id='organization-structure'):
    return {
        'schema_version': 1, 'config_id': config_id, 'version': 1,
        'company': {'en': company, 'ar': 'الشركة التجريبية'},
        'business_units': [
            {'key': 'group', 'label': {'en': 'Group', 'ar': 'المجموعة'}},
            {'key': 'markets', 'parent': 'group', 'label': {'en': 'Markets', 'ar': 'الأسواق'}},
        ],
        'dimensions': [
            {'key': 'region', 'label': {'en': 'Region', 'ar': 'المنطقة'}, 'members': [
                {'key': 'all-regions', 'label': {'en': 'All regions', 'ar': 'كل المناطق'}},
                {'key': 'north', 'parent': 'all-regions', 'label': {'en': 'North', 'ar': 'الشمال'}},
            ]},
            {'key': 'client', 'label': {'en': 'Client', 'ar': 'العميل'}, 'members': [
                {'key': 'all-clients', 'label': {'en': 'All clients', 'ar': 'كل العملاء'}},
                {'key': 'institution', 'parent': 'all-clients',
                 'label': {'en': 'Institution', 'ar': 'مؤسسة'}},
            ]},
        ],
        'source_mappings': [
            {'source_key': 'intent-proof', 'source_field': 'business_unit',
             'target_type': 'business_unit', 'values': [
                 {'source_value': 'Group', 'target': 'group'},
                 {'source_value': 'Markets', 'target': 'markets'}]},
            {'source_key': 'intent-proof', 'source_field': 'region',
             'target_type': 'dimension', 'target_key': 'region', 'values': [
                 {'source_value': 'ALL', 'target': 'all-regions'},
                 {'source_value': 'N', 'target': 'north'}]},
            {'source_key': 'intent-proof', 'source_field': 'counterparty',
             'target_type': 'dimension', 'target_key': 'client', 'values': [
                 {'source_value': 'ALL', 'target': 'all-clients'},
                 {'source_value': 'INST', 'target': 'institution'}]},
        ],
    }


def fixture_structure_body():
    dimensions = []
    mappings = [{
        'source_key': 'intent-proof', 'source_field': 'business_unit',
        'target_type': 'business_unit',
        'values': [{'source_value': 'Group', 'target': 'group'}],
    }]
    configured = {
        'product': ['item-a', 'item-b'],
        'region': ['all-regions', 'north', 'south'],
        'client': ['retail', 'institution', 'hospital', 'pharmacy'],
    }
    for key, members in configured.items():
        dimension_members = [
            {'key': member, 'label': {'en': member.title(), 'ar': 'قيمة ' + member},
             **({'parent': 'all-regions'} if key == 'region' and member != 'all-regions' else {})}
            for member in members
        ]
        dimensions.append({
            'key': key, 'label': {'en': key.title(), 'ar': 'بُعد ' + key},
            'members': dimension_members,
        })
        mappings.append({
            'source_key': 'intent-proof', 'source_field': key,
            'target_type': 'dimension', 'target_key': key,
            'values': [{'source_value': member.upper(), 'target': member} for member in members],
        })
    return {
        'schema_version': 1, 'config_id': 'fixture-structure', 'version': 1,
        'company': {'en': 'Fixture Company', 'ar': 'شركة الاختبار'},
        'business_units': [{'key': 'group', 'label': {'en': 'Group', 'ar': 'المجموعة'}}],
        'dimensions': dimensions, 'source_mappings': mappings,
    }


def test_tenant_structure_contract_is_sector_neutral_hierarchical_and_complete():
    healthcare = TenantStructureConfiguration.model_validate(structure_body('Healthcare Distribution'))
    exchange_body = structure_body('Securities Exchange', 'exchange-structure')
    exchange_body['dimensions'] = [{
        'key': 'instrument', 'label': {'en': 'Instrument', 'ar': 'الأداة'},
        'members': [{'key': 'equity', 'label': {'en': 'Equity', 'ar': 'أسهم'}}],
    }]
    exchange_body['source_mappings'] = [exchange_body['source_mappings'][0], {
        'source_key': 'intent-proof', 'source_field': 'security_type',
        'target_type': 'dimension', 'target_key': 'instrument',
        'values': [{'source_value': 'EQ', 'target': 'equity'}],
    }]
    exchange = TenantStructureConfiguration.model_validate(exchange_body)
    assert healthcare.dimensions[0].key == 'region'
    assert exchange.dimensions[0].key == 'instrument'
    broken = structure_body()
    broken['business_units'][0]['parent'] = 'markets'
    with pytest.raises(ValueError, match='cycle'):
        TenantStructureConfiguration.model_validate(broken)
    broken = structure_body()
    broken['source_mappings'].pop()
    with pytest.raises(ValueError, match='required'):
        TenantStructureConfiguration.model_validate(broken)
    broken = structure_body()
    broken['source_mappings'][1]['values'][1]['target'] = 'unconfigured'
    with pytest.raises(ValueError, match='not configured'):
        TenantStructureConfiguration.model_validate(broken)


def test_tenant_structure_versions_readiness_approval_and_api(setup):
    s = setup
    configuration = TenantStructureConfiguration.model_validate(structure_body())
    created = tenant_structure_store.create(s['operator'], configuration)
    assert created['readiness']['status'] == 'ready'
    assert all(created['readiness']['checks'].values())
    assert created['source_bindings'] == [{'source_key': 'intent-proof', 'source_system_id': 1}]
    assert tenant_structure_store.create(s['operator'], configuration) == created
    with pytest.raises(PermissionError):
        tenant_structure_store.approve(s['operator'], configuration.config_id, 1, created['digest'],
                                       'Reviewed organization, dimensions and mappings.')
    approval = tenant_structure_store.approve(s['admin'], configuration.config_id, 1, created['digest'],
                                               'Reviewed organization, dimensions and mappings independently.')
    assert approval['config_digest'] == created['digest']
    approved = tenant_structure_store.read(s['executive'], configuration.config_id, 1)
    assert approved['readiness']['status'] == 'approved'
    assert approved['authoritative'] is True
    changed = structure_body(); changed['company']['en'] = 'Changed'
    with pytest.raises(store.Conflict, match='different content'):
        tenant_structure_store.create(s['operator'], TenantStructureConfiguration.model_validate(changed))
    second = structure_body(); second['version'] = 2; second['company']['en'] = 'Example Holdings v2'
    proposed = tenant_structure_store.create(s['operator'], TenantStructureConfiguration.model_validate(second))
    assert proposed['authoritative'] is False
    assert tenant_structure_store.read(s['executive'], configuration.config_id, 1)['authoritative'] is True
    skipped = structure_body(); skipped['version'] = 4
    with pytest.raises(store.Conflict, match='consecutive'):
        tenant_structure_store.create(s['operator'], TenantStructureConfiguration.model_validate(skipped))
    unknown = structure_body(config_id='unknown-source')
    unknown['source_mappings'][0]['source_key'] = 'not-registered'
    with pytest.raises(ValueError, match='Register'):
        tenant_structure_store.create(s['operator'], TenantStructureConfiguration.model_validate(unknown))
    s['current']['principal'] = s['operator']
    listed = s['client'].get('/api/intent/dimensional/advisor/structure-configurations')
    assert listed.status_code == 200
    assert any(item['config_id'] == 'organization-structure' and item['version'] == 1 and item['approved']
               for item in listed.json()['configurations'])
    read = s['client'].get('/api/intent/dimensional/advisor/structure-configurations/organization-structure/versions/1')
    assert read.status_code == 200 and read.json()['payload']['dimensions'][0]['key'] == 'region'
    assert s['client'].post('/api/intent/dimensional/advisor/structure-configurations',
                            json=configuration.model_dump(mode='json')).status_code == 200
    with s['connect']() as conn:
        for table in ['strategyos_intent_structure_configs', 'strategyos_intent_structure_approvals']:
            with pytest.raises(psycopg.Error, match='immutable'):
                conn.execute(sql.SQL('UPDATE {} SET tenant_key=tenant_key').format(sql.Identifier(table)))
            conn.rollback()


def test_plan_import_and_ratification_require_current_approved_structure(setup):
    s = setup
    unbound = deepcopy(s['p'])
    unbound.pop('business_unit'); unbound.pop('structure')
    with pytest.raises(ValueError, match='approved organization-structure'):
        store.import_plan(s['operator'], Plan.model_validate(unbound), s['pack'])
    changed_digest = deepcopy(s['p'])
    changed_digest['structure']['digest'] = '0' * 64
    with pytest.raises(store.Conflict, match='fingerprint'):
        store.import_plan(s['operator'], Plan.model_validate(changed_digest), s['pack'])
    unknown_member = deepcopy(s['p'])
    unknown_member['dimensions']['client'].append('unconfigured-client')
    with pytest.raises(ValueError, match='outside the approved'):
        store.import_plan(s['operator'], Plan.model_validate(unknown_member), s['pack'])
    imported = store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    current = store.read_plan(s['executive'], s['p']['plan_id'], 1)
    assert current['structure']['status'] == 'current'
    assert current['structure']['business_unit'] == 'group'

    successor_body = fixture_structure_body()
    successor_body['version'] = 2
    successor_body['company']['en'] = 'Fixture Company Updated'
    successor = tenant_structure_store.create(
        s['operator'], TenantStructureConfiguration.model_validate(successor_body))
    pending_structure = deepcopy(s['p'])
    pending_structure['version'] = 2
    pending_structure['structure'] = {'config_id': 'fixture-structure', 'version': 2,
                                      'digest': successor['digest']}
    with pytest.raises(store.Conflict, match='independently approved'):
        store.import_plan(s['operator'], Plan.model_validate(pending_structure), s['pack'])
    tenant_structure_store.approve(
        s['admin'], 'fixture-structure', 2, successor['digest'],
        'Independently reviewed the updated fixture organization structure.')
    assert store.read_plan(s['executive'], s['p']['plan_id'], 1)['structure']['status'] == 'superseded'
    store.set_ratifier(s['admin'], s['p']['plan_id'], s['executive']['subject'], True, 0)
    with pytest.raises(store.Conflict, match='superseded'):
        store.ratify(s['executive'], s['p']['plan_id'], 1, imported['digest'],
                     'Reviewed the targets, owners and source evidence.')


def test_durable_roundtrip_and_immutable_history(setup):
    s = setup
    imported = import_pair(s)
    assert store.read_plan(s['executive'], s['p']['plan_id'], 1)['governance_status'] == 'proposed'
    with pytest.raises(store.Conflict, match='not ratified'): analyse(s)
    approval = approve(s, imported)
    result = analyse(s)
    assert result['rollups'][0]['offset_detected']
    assert result['approval_status'] == 'ratified'
    assert result['ratification']['approved_by'] == 'executive'
    assert store.read_plan(s['executive'], s['p']['plan_id'], 1)['ratification']['plan_digest'] == imported['digest']
    assert store.ratify(s['executive'], s['p']['plan_id'], 1, imported['digest'], approval['note']) == approval
    assert analyse(s) == result
    assert store.read_analysis(s['executive'], result['analysis_hash']) == result
    # Every read uses a fresh connection; there is no process-memory store.
    (s['root'] / 'raw' / 'evidence.csv').write_text('later source change')
    assert store.read_analysis(s['executive'], result['analysis_hash']) == result
    with pytest.raises(sources.SourceUnavailable): analyse(s)
    with s['connect']() as conn:
        for table in ['strategyos_intent_plan_versions', 'strategyos_intent_actual_versions',
                      'strategyos_intent_ratifier_events', 'strategyos_intent_ratifications', 'strategyos_intent_analyses']:
            for action in ['DELETE FROM {}', 'UPDATE {} SET tenant_key=tenant_key', 'TRUNCATE {} CASCADE']:
                with pytest.raises(psycopg.Error, match='immutable'):
                    conn.execute(sql.SQL(action).format(sql.Identifier(table)))
                conn.rollback()


def test_concurrent_import_and_ratification_retries(setup):
    s = setup
    with ThreadPoolExecutor(max_workers=4) as pool:
        plans = list(pool.map(lambda _: store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack']), range(4)))
    assert all(p == plans[0] for p in plans)
    store.set_ratifier(s['admin'], s['p']['plan_id'], 'executive', True, 0)
    with ThreadPoolExecutor(max_workers=4) as pool:
        approvals = list(pool.map(lambda _: store.ratify(s['executive'], s['p']['plan_id'], 1, plans[0]['digest'],
                                                        'Reviewed targets and sources.'), range(4)))
    assert all(a == approvals[0] for a in approvals)
    with s['connect']() as conn:
        assert conn.execute('SELECT count(*) FROM strategyos_intent_ratifications').fetchone()[0] == 1


def test_versions_conflicts_supersession_and_historical_results(setup):
    s = setup
    plan = import_pair(s)
    approve(s, plan)
    old = analyse(s)
    s['p']['cells'][0]['owner'] = 'changed-owner'
    with pytest.raises(store.Conflict): store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    s['p']['version'] = 3
    with pytest.raises(store.Conflict): store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    s['p']['version'] = 2
    second = store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    assert second['digest'] != plan['digest']
    store.ratify(s['executive'], s['p']['plan_id'], 2, second['digest'], 'Reviewed the changed owner and targets.')
    with pytest.raises(store.Conflict, match='newer ratified'):
        store.create_analysis(s['executive'], s['p']['plan_id'], 1, s['a']['revision'], TODAY)
    assert store.read_analysis(s['executive'], old['analysis_hash']) == old
    assert analyse(s)['plan_version'] == 2
    s['a']['observations'][0]['value'] = '41'
    with pytest.raises(store.Conflict): store.import_actuals(s['operator'], Actuals.model_validate(s['a']), s['pack'])
    s['a']['revision'] = 'revision-2'
    store.import_actuals(s['operator'], Actuals.model_validate(s['a']), s['pack'])
    assert analyse(s)['rollups'][0]['actual'] == '201'


def test_grant_required_separation_of_duties_and_revocation(setup):
    s = setup
    plan = store.import_plan(s['admin'], Plan.model_validate(s['p']), s['pack'])
    with pytest.raises(PermissionError): store.ratify(s['executive'], s['p']['plan_id'], 1, plan['digest'], 'Reviewed all source targets.')
    with pytest.raises(PermissionError): store.set_ratifier(s['admin'], s['p']['plan_id'], 'admin', True, 0)
    other_admin = {**s['admin'], 'subject': 'other-admin'}
    store.set_ratifier(other_admin, s['p']['plan_id'], 'admin', True, 0)
    with pytest.raises(PermissionError): store.ratify(s['admin'], s['p']['plan_id'], 1, plan['digest'], 'Reviewed all source targets.')
    store.set_ratifier(s['admin'], s['p']['plan_id'], 'executive', True, 0)
    with pytest.raises(store.Conflict): store.set_ratifier(s['admin'], s['p']['plan_id'], 'executive', False, 0)
    assert store.read_ratifier(s['admin'], s['p']['plan_id'], 'executive')['revision'] == 1
    store.set_ratifier(s['admin'], s['p']['plan_id'], 'executive', False, 1)
    with pytest.raises(PermissionError): store.ratify(s['executive'], s['p']['plan_id'], 1, plan['digest'], 'Reviewed all source targets.')
    with s['connect']() as conn:
        assert conn.execute('SELECT count(*) FROM strategyos_intent_ratifications').fetchone()[0] == 0


@pytest.mark.parametrize('principal_change', [{'tenant_id': 'other'}, {'role': 'bu'}, {'role': 'anonymous'},
                                             {'auth_disabled': True}, {'demo_role_login': True}])
def test_scope_denial_before_storage_or_sources(setup, principal_change, monkeypatch):
    s = setup
    principal = {**s['executive'], **principal_change}
    monkeypatch.setattr(store.state_store, 'database_connection', lambda: pytest.fail('Accessed store before authorization'))
    with pytest.raises(PermissionError): store.read_plan(principal, s['p']['plan_id'], 1)
    with pytest.raises(PermissionError): store.read_analysis(principal, 'a' * 64)


@pytest.mark.parametrize('disposition', ['restricted_context', 'evaluator_only', 'control_plane', 'quarantined_context'])
def test_nonbusiness_sources_cannot_be_imported_or_read(setup, disposition):
    s = setup
    imported = import_pair(s)
    approve(s, imported)
    result = analyse(s)
    s['manifest']['manifest'][0]['source_disposition'] = disposition
    (s['root'] / 'summary.json').write_text(json.dumps(s['manifest']))
    with pytest.raises(sources.SourceUnavailable): store.read_analysis(s['executive'], result['analysis_hash'])
    with pytest.raises(sources.SourceUnavailable): store.read_plan(s['executive'], s['p']['plan_id'], 1)
    with pytest.raises(sources.SourceUnavailable): store.import_actuals(s['operator'], Actuals.model_validate(s['a']), s['pack'])


def test_foreign_pack_and_spoofed_approval_rejected(setup):
    s = setup
    s['p'].update(status='ratified', ratified_by='executive', ratified_on=s['p']['period']['start'],
                  ratification=s['p']['cells'][0]['source'])
    with pytest.raises(ValueError, match='proposals'): store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    s['manifest']['tenant_context']['tenant_id'] = 'foreign'
    (s['root'] / 'summary.json').write_text(json.dumps(s['manifest']))
    with pytest.raises(PermissionError): store.import_actuals(s['operator'], Actuals.model_validate(s['a']), s['pack'])


def test_missing_database_does_not_fall_back_to_files(setup, monkeypatch):
    s = setup
    monkeypatch.setattr(store.state_store, 'database_connection', lambda: (None, 'not configured'))
    with pytest.raises(store.Unavailable): store.read_plan(s['executive'], 'missing', 1)


def test_two_source_packs_with_same_filename_are_verified_independently(setup):
    s = setup
    other = s['root'].parent / 'actual-pack'
    shutil.copytree(s['root'], other)
    import hashlib
    (other / 'raw' / 'evidence.csv').write_text('separate actual evidence')
    sha = hashlib.sha256((other / 'raw' / 'evidence.csv').read_bytes()).hexdigest()
    manifest = deepcopy(s['manifest'])
    manifest['source_pack_id'] = 'actual-pack'
    manifest['manifest'][0]['sha256'] = sha
    (other / 'summary.json').write_text(json.dumps(manifest))
    for obs in s['a']['observations']: obs['source']['sha256'] = sha
    plan = store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    store.import_actuals(s['operator'], Actuals.model_validate(s['a']), 'actual-pack')
    approve(s, plan)
    result = analyse(s)
    assert result['source_receipts']['actuals']['source_pack_id'] == 'actual-pack'
    assert result['rollups'][0]['actual'] == '200'
    (other / 'raw' / 'evidence.csv').write_text('changed')
    with pytest.raises(sources.SourceUnavailable): analyse(s)


def test_api_real_workflow_and_no_cache(setup):
    s = setup
    client = s['client']
    prefix = '/api/intent/dimensional'
    response = client.post(prefix + '/plans', json={'source_pack_id': s['pack'], 'plan': s['p']})
    assert response.status_code == 200, response.text
    assert response.headers['cache-control'] == 'private, no-store'
    digest = response.json()['digest']
    response = client.post(prefix + '/actuals', json={'source_pack_id': s['pack'], 'actuals': s['a']})
    assert response.status_code == 200, response.text
    path = prefix + '/plans/' + s['p']['plan_id']
    s['current']['principal'] = s['admin']
    response = client.put(path + '/ratifier', json={'subject': 'executive', 'enabled': True, 'expected_revision': 0})
    assert response.status_code == 200, response.text
    s['current']['principal'] = s['executive']
    assert client.post(path + '/versions/1/ratify', json={'expected_digest': 'a'*64, 'note': 'Reviewed targets and owners.'}).status_code == 409
    response = client.post(path + '/versions/1/ratify', json={'expected_digest': digest, 'note': 'Reviewed targets and owners.'})
    assert response.status_code == 200, response.text
    response = client.post(prefix + '/analyses', json={'plan_id': s['p']['plan_id'], 'plan_version': 1,
                                                       'actual_revision': s['a']['revision'], 'as_of': TODAY.isoformat()})
    assert response.status_code == 200, response.text
    result = response.json()
    assert client.get(prefix + '/analyses/' + result['analysis_hash']).json() == result
    assert client.get(path + '/versions/1').json()['governance_status'] == 'ratified'
    assert client.get(prefix + '/actuals/' + s['a']['revision']).status_code == 200
    assert client.get(path + '/versions/99').status_code == 404


def test_api_permission_payload_and_csrf_boundaries(setup, monkeypatch):
    s, prefix = setup, '/api/intent/dimensional'
    client = s['client']
    s['current']['principal'] = s['executive']
    assert client.post(prefix + '/plans', json={'source_pack_id': s['pack'], 'plan': s['p']}).status_code == 403
    s['current']['principal'] = s['operator']
    assert client.post(prefix + '/plans', json={'source_pack_id': s['pack'], 'plan': s['p'], 'role': 'tenant_admin'}).status_code == 422
    assert client.post(prefix + '/plans', content=b'x' * (store.MAX_BYTES + 1)).status_code == 413
    client.cookies.set('strategyos_session', 'test-cookie')
    assert client.post(prefix + '/plans', headers={'Origin': 'https://foreign.invalid'},
                       json={'source_pack_id': s['pack'], 'plan': s['p']}).status_code == 403
    assert client.post(prefix + '/plans', headers={'Origin': 'http://testserver'},
                       json={'source_pack_id': s['pack'], 'plan': s['p']}).status_code == 200
    client.cookies.clear()
    monkeypatch.setattr(store.state_store, 'database_connection', lambda: (None, 'offline'))
    assert client.get(prefix + '/plans/anything/versions/1').status_code == 503


def test_api_uses_real_authentication_and_registered_app_routes(setup, monkeypatch):
    s = setup
    from strategyos_mvp import api
    config = replace(auth.CONFIG, api_auth_enabled=True, auth_mode='api_key', idp_enabled=False,
                     demo_role_login_enabled=False, tenant_slug='intent-tenant',
                     operator_api_keys=('local-intent-test-key',))
    monkeypatch.setattr(auth, 'CONFIG', config)
    # Exercise the real app's identity middleware and registered router, with no dependency overrides.
    client = TestClient(api.app)
    prefix = '/api/intent/dimensional'
    assert client.get(prefix + '/plans/unknown/versions/1').status_code == 401
    assert client.post(prefix + '/plans', json={'source_pack_id': s['pack'], 'plan': s['p']}).status_code == 401
    response = client.post(prefix + '/plans', headers={'X-API-Key': 'local-intent-test-key'},
                           json={'source_pack_id': s['pack'], 'plan': s['p']})
    assert response.status_code == 200, response.text
    assert response.json()['imported_by'].startswith('api-key:operator:')
    assert response.headers['cache-control'] == 'private, no-store'


def test_period_overlap_and_unaddressable_identifiers_rejected(setup):
    s = setup
    import_pair(s)
    s['p']['version'] = 2
    s['p']['period']['start'] = f'{TODAY.year-1}-01-15'
    with pytest.raises(store.Conflict, match='Overlapping'):
        store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    s['p']['plan_id'] = 'path/segment'
    with pytest.raises(ValueError, match='identifier'):
        store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])


def test_protected_evidence_download_checks_bytes_and_never_accepts_paths(setup):
    s = setup
    approve(s, import_pair(s))
    result = analyse(s)
    path = '/api/intent/dimensional/analyses/' + result['analysis_hash'] + '/evidence'
    response = s['client'].get(path, params={'cell_id': 'regional', 'side': 'actuals'})
    assert response.status_code == 200, response.text
    assert response.content == (s['root'] / 'raw' / 'evidence.csv').read_bytes()
    assert response.headers['x-content-type-options'] == 'nosniff'
    assert s['client'].get(path, params={'cell_id': '../evidence.csv', 'side': 'plan'}).status_code == 404
    (s['root'] / 'raw' / 'evidence.csv').write_text('new bytes')
    assert s['client'].get(path, params={'cell_id': 'regional', 'side': 'actuals'}).status_code == 409


def test_symlink_and_forged_evaluator_manifest_rejected(setup):
    s = setup
    raw = s['root'] / 'raw'
    (raw / 'hidden.csv').write_bytes((raw / 'evidence.csv').read_bytes())
    (raw / 'evidence.csv').unlink()
    (raw / 'evidence.csv').symlink_to(raw / 'hidden.csv')
    with pytest.raises(sources.SourceUnavailable, match='symbolic'):
        store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    (raw / 'evidence.csv').unlink()
    shutil.copy(raw / 'hidden.csv', raw / 'answer_key.csv')
    for cell in s['p']['cells']: cell['source']['path'] = 'answer_key.csv'
    s['manifest']['manifest'][0]['relative_path'] = 'answer_key.csv'
    (s['root'] / 'summary.json').write_text(json.dumps(s['manifest']))
    with pytest.raises(sources.SourceUnavailable):
        store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])


def test_shared_database_cannot_reveal_another_tenants_records(setup, monkeypatch):
    s = setup
    approve(s, import_pair(s))
    saved = analyse(s)
    monkeypatch.setattr(store, 'CONFIG', replace(store.CONFIG, tenant_slug='second-tenant'))
    other = {**s['executive'], 'tenant_id': 'second-tenant'}
    with pytest.raises(store.NotFound): store.read_plan(other, s['p']['plan_id'], 1)
    with pytest.raises(store.NotFound): store.read_actuals(other, s['a']['revision'])
    with pytest.raises(store.NotFound): store.read_analysis(other, saved['analysis_hash'])


def test_migration_is_idempotent_and_not_run_implicitly(setup):
    s = setup
    import_pair(s)
    store.initialize()
    assert store.read_plan(s['executive'], s['p']['plan_id'], 1)['governance_status'] == 'proposed'
    with s['connect']() as conn:
        conn.execute('DROP TABLE strategyos_intent_analyses CASCADE')
    with pytest.raises(store.Unavailable, match='migration'):
        store.read_plan(s['executive'], s['p']['plan_id'], 1)


def test_as_of_cannot_precede_the_server_import_or_ratification(setup):
    s = setup
    approve(s, import_pair(s))
    from datetime import timedelta
    with pytest.raises(store.Conflict, match='imported after'):
        store.create_analysis(s['executive'], s['p']['plan_id'], 1, s['a']['revision'], TODAY - timedelta(days=1))
    with pytest.raises(ValueError, match='future'):
        store.create_analysis(s['executive'], s['p']['plan_id'], 1, s['a']['revision'], TODAY + timedelta(days=1))


def test_catalog_permissions_pagination_and_source_revocation(setup):
    s = setup
    assert store.catalog(s['executive'])['plans'] == []
    imported = import_pair(s)
    approve(s, imported)
    s['p']['version'] = 2
    store.import_plan(s['operator'], Plan.model_validate(s['p']), s['pack'])
    page = store.catalog(s['executive'], limit=1)
    assert page['plans'][0]['version'] == 2
    assert page['permissions'] == {'can_import': False, 'can_manage_ratifiers': False}
    assert page['next_offset'] == 1
    assert store.catalog(s['executive'], offset=1, limit=1)['plans'][0]['version'] == 1
    assert store.catalog(s['admin'])['permissions'] == {'can_import': True, 'can_manage_ratifiers': True}
    assert store.read_plan(s['executive'], s['p']['plan_id'], 2)['permissions']['can_ratify'] is True
    assert store.read_plan(s['operator'], s['p']['plan_id'], 2)['permissions']['can_ratify'] is False
    s['manifest']['manifest'][0]['source_disposition'] = 'restricted_context'
    (s['root'] / 'summary.json').write_text(json.dumps(s['manifest']))
    catalog = store.catalog(s['executive'])
    assert catalog['plans'] == catalog['actuals'] == []


def test_quality_assurance_plans_require_explicit_catalog_marker(setup):
    s = setup
    qa = deepcopy(s['p'])
    qa['plan_id'] = 'qa-plan-visible-only-to-test'
    qa['catalog_visibility'] = 'quality_assurance'
    store.import_plan(s['operator'], Plan.model_validate(qa), s['pack'])

    assert store.catalog(s['operator'])['plans'] == []
    visible = store.catalog(s['operator'], qa_plan_id=qa['plan_id'])['plans']
    assert [(item['plan_id'], item['version']) for item in visible] == [(qa['plan_id'], 1)]


def test_legacy_hosted_test_plans_are_not_customer_catalog_records(setup):
    s = setup
    qa = deepcopy(s['p'])
    qa['plan_id'] = 'human-plan-20260909123456'
    store.import_plan(s['operator'], Plan.model_validate(qa), s['pack'])

    assert store.catalog(s['operator'])['plans'] == []


def test_reviewer_can_open_proposal_evidence_before_ratification(setup):
    s = setup
    import_pair(s)
    path = '/api/intent/dimensional/plans/' + s['p']['plan_id'] + '/versions/1/evidence'
    s['current']['principal'] = s['executive']
    result = s['client'].get(path, params={'cell_id': 'regional'})
    assert result.status_code == 200
    assert result.content == (s['root'] / 'raw' / 'evidence.csv').read_bytes()


def test_live_source_policy_revocation_and_export_are_enforced(setup):
    s=setup
    plan=import_pair(s)
    with s['connect']() as conn:
        conn.execute("UPDATE strategyos_source_access_policies SET export_allowed=false")
    store.read_plan(s['operator'], s['p']['plan_id'], 1)
    with pytest.raises(PermissionError):
        store.plan_evidence_bytes(s['operator'], s['p']['plan_id'], 1, s['p']['cells'][0]['id'])
    with s['connect']() as conn:
        conn.execute("UPDATE strategyos_source_access_policies SET allowed_roles=ARRAY[]::text[]")
    with pytest.raises(PermissionError):
        store.read_plan(s['operator'], s['p']['plan_id'], 1)
    assert store.catalog(s['operator'])['plans']==[]


def test_intent_rls_blocks_unscoped_and_other_tenant_reads(setup):
    s=setup
    import_pair(s)
    role='intent_reader_'+uuid4().hex
    with s['connect']() as conn:
        conn.execute(sql.SQL('CREATE ROLE {} NOLOGIN').format(sql.Identifier(role)))
        conn.execute(sql.SQL('GRANT SELECT ON strategyos_intent_plan_versions TO {}').format(sql.Identifier(role)))
        conn.execute(sql.SQL('SET ROLE {}').format(sql.Identifier(role)))
        assert conn.execute('SELECT count(*) FROM strategyos_intent_plan_versions').fetchone()[0]==0
        conn.execute("SELECT set_config('strategyos.tenant_key','another-tenant',true)")
        assert conn.execute('SELECT count(*) FROM strategyos_intent_plan_versions').fetchone()[0]==0
        conn.execute("SELECT set_config('strategyos.tenant_key','intent-tenant',true)")
        assert conn.execute('SELECT count(*) FROM strategyos_intent_plan_versions').fetchone()[0]==1
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            conn.execute('DELETE FROM strategyos_intent_plan_versions')


def test_cookie_writes_use_configured_https_origin_behind_proxy(setup, monkeypatch):
    s=setup
    monkeypatch.setenv('STRATEGYOS_PUBLIC_URL','https://preview.example.test/')
    s['client'].cookies.set('strategyos_session','test-cookie')
    body={'source_pack_id':s['pack'],'plan':s['p']}
    for headers in ({'Origin':'http://testserver'}, {'Origin':'https://attacker.test','X-Forwarded-Host':'attacker.test','X-Forwarded-Proto':'https'}):
        assert s['client'].post('/api/intent/dimensional/plans',headers=headers,json=body).status_code==403
    assert s['client'].post('/api/intent/dimensional/plans',headers={'Origin':'https://preview.example.test'},json=body).status_code==200


def test_board_pack_exports_bound_snapshot_and_translations(setup, tmp_path, monkeypatch):
    from io import BytesIO
    from pptx import Presentation
    from pypdf import PdfReader
    from strategyos_mvp import board_pack
    s = setup
    approve(s, import_pair(s))
    result = analyse(s)
    request = board_pack.PackRequest(template=board_pack.PackTemplate(labels={
        'revenue': board_pack.Translation(en='Revenue', ar='الإيرادات')}))
    pack = board_pack.compose(s['executive'], result['analysis_hash'], request)
    assert pack['binding']['analysis_hash'] == result['analysis_hash']
    assert not pack['binding']['warnings']
    assert len(pack['evidence']) == len(result['cells']) * 2
    assert any('الإيرادات' in line for p in pack['pages'] for line in p['lines'])
    assert any('masks' in line for p in pack['pages'] for line in p['lines'])
    for c in result['cells']:
        page = next(p for p in pack['pages'] if p['lines'][0].startswith(c['cell_id'] + ' ·'))
        assert page['lines'][3].endswith(': ' + c['target'])
        assert page['lines'][4].endswith(': ' + c['actual'])
    assert board_pack.compose(s['executive'], result['analysis_hash'], request)['pack_hash'] == pack['pack_hash']
    pdf = board_pack.export_pdf(pack, 'https://kyvern.example')
    pptx = board_pack.export_pptx(pack, 'https://kyvern.example')
    assert len(PdfReader(BytesIO(pdf)).pages) == len(pack['pages'])
    deck = Presentation(BytesIO(pptx))
    assert len(deck.slides) == len(pack['pages'])
    assert result['analysis_hash'] in deck.slides[0].notes_slide.notes_text_frame.text
    assert any('https://kyvern.example/api/intent/' in str(r.target_ref) for slide in deck.slides for r in slide.part.rels.values())
    # Local artifacts are optional and never created in CI or used as test truth.
    import os
    if os.environ.get('KYVERN_PACK_QA_DIR'):
        root = Path(os.environ['KYVERN_PACK_QA_DIR']); root.mkdir(parents=True, exist_ok=True)
        (root / 'board-bilingual.pdf').write_bytes(pdf)
        (root / 'board-bilingual.pptx').write_bytes(pptx)
        (root / 'board.json').write_text(json.dumps(pack, ensure_ascii=False, indent=2))
    s['current']['principal'] = s['executive']
    url = '/api/intent/dimensional/analyses/' + result['analysis_hash'] + '/board-pack'
    assert s['client'].post(url, json=request.model_dump()).status_code == 200
    for format in ['pdf', 'pptx']:
        response = s['client'].post(url + '/' + format, json=request.model_dump())
        assert response.status_code == 200
        assert response.headers['Cache-Control'] == 'private, no-store'
        assert response.headers['X-Kyvern-Pack-Hash'] == pack['pack_hash']
    assert s['client'].post(url, json={'template': {'numbers': {'revenue': '999'}}}).status_code == 422
    assert s['client'].post(url, json={'language': 'xx'}).status_code == 422
    s['client'].cookies.set('strategyos_session', 'test-cookie')
    monkeypatch.setenv('STRATEGYOS_PUBLIC_URL', 'https://kyvern.example')
    assert s['client'].post(url + '/pdf', json={}).status_code == 403
    assert s['client'].post(url + '/pdf', json={}, headers={'Origin': 'https://kyvern.example'}).status_code == 200


def test_board_pack_revocation_integrity_and_tenant_boundaries(setup):
    from strategyos_mvp import board_pack
    s = setup
    approve(s, import_pair(s)); result = analyse(s)
    req = board_pack.PackRequest()
    for role in ['bu', 'system', 'analyst']:
        with pytest.raises(PermissionError):
            board_pack.compose({**s['executive'], 'role': role}, result['analysis_hash'], req)
    with pytest.raises(PermissionError):
        board_pack.compose({**s['executive'], 'tenant_id': 'other'}, result['analysis_hash'], req)
    with s['connect']() as conn:
        conn.execute('UPDATE strategyos_source_access_policies SET export_allowed=false')
    assert store.read_analysis(s['executive'], result['analysis_hash']) == result
    with pytest.raises(PermissionError): board_pack.compose(s['executive'], result['analysis_hash'], req)
    with s['connect']() as conn:
        conn.execute('UPDATE strategyos_source_access_policies SET export_allowed=true')
    (s['root'] / 'raw' / 'evidence.csv').write_text('changed evidence')
    with pytest.raises(sources.SourceUnavailable): board_pack.compose(s['executive'], result['analysis_hash'], req)


def test_board_pack_freshness_does_not_replace_saved_figures(setup):
    from strategyos_mvp import board_pack
    s = setup
    approve(s, import_pair(s)); result = analyse(s)
    original = board_pack.compose(s['executive'], result['analysis_hash'], board_pack.PackRequest())
    newer = deepcopy(s['a']); newer['revision'] = 'newer-actuals'
    store.import_actuals(s['operator'], Actuals.model_validate(newer), s['pack'])
    proposal = deepcopy(s['p']); proposal['version'] = 2
    imported = store.import_plan(s['operator'], Plan.model_validate(proposal), s['pack'])
    # A proposal is not a new board-approved comparator.
    fresh = board_pack.compose(s['executive'], result['analysis_hash'], board_pack.PackRequest())
    assert fresh['binding']['warnings'] == ['new_actuals']
    store.ratify(s['executive'], proposal['plan_id'], 2, imported['digest'], 'Reviewed the next plan version and its evidence.')
    fresh = board_pack.compose(s['executive'], result['analysis_hash'], board_pack.PackRequest())
    assert fresh['binding']['warnings'] == ['new_plan', 'new_actuals']
    assert fresh['binding']['actual_digest'] == original['binding']['actual_digest']
    assert fresh['binding']['plan_digest'] == original['binding']['plan_digest']
    assert store.read_analysis(s['executive'], result['analysis_hash']) == result


def test_board_pack_missing_actuals_stay_missing(setup):
    from strategyos_mvp import board_pack
    s = setup; s['a']['observations'].pop()
    approve(s, import_pair(s)); result = analyse(s)
    pack = board_pack.compose(s['executive'], result['analysis_hash'], board_pack.PackRequest(language='en'))
    assert any(line == 'Actual: Missing' for p in pack['pages'] for line in p['lines'])
    assert any(line == 'Status: Incomplete' for p in pack['pages'] for line in p['lines'])


def test_board_pack_bidi_preserves_signed_decimals_and_dates():
    from strategyos_mvp.board_pack import visual, mark_ltr
    for value in ['2026-09-07', '-60.25', '+80.50', '0.0001']:
        assert value in visual('الانحراف: ' + value)
        assert '\u200e' + value + '\u200e' in mark_ltr('الانحراف: ' + value)


def test_recorded_board_template_and_pack_history_are_immutable_and_downloadable(setup):
    from strategyos_mvp import board_pack, board_pack_store
    s = setup
    approve(s, import_pair(s)); analysis = analyse(s)
    template = board_pack.PackTemplate(
        template_id='healthcare-board', version=1,
        labels={'revenue': board_pack.Translation(en='Revenue', ar='الإيرادات')})
    registered = board_pack_store.register(s['operator'], template)
    assert registered['digest'] == fingerprint(template.model_dump(mode='json'))
    assert registered['origin'] == {'type': 'manual'}
    assert board_pack_store.register(s['operator'], template) == registered
    changed = template.model_copy(update={'accent': '#123456'})
    with pytest.raises(store.Conflict, match='different content'):
        board_pack_store.register(s['operator'], changed)
    with pytest.raises(store.Conflict, match='consecutive'):
        board_pack_store.register(s['operator'], template.model_copy(update={'version': 3}))
    record = board_pack_store.create_pack(s['executive'], analysis['analysis_hash'],
                                          template.template_id, template.version, 'bilingual')
    assert record['freshness']['status'] == 'current'
    assert record['pack_id'] == record['pack']['pack_hash']
    assert record['pack']['binding']['analysis_hash'] == analysis['analysis_hash']
    repeated = board_pack_store.create_pack(s['executive'], analysis['analysis_hash'],
                                            template.template_id, template.version, 'bilingual')
    assert repeated['pack'] == record['pack']
    assert repeated['digest'] == record['digest']
    history = board_pack_store.packs(s['executive'], analysis['analysis_hash'])
    assert [item['pack_id'] for item in history['packs']] == [record['pack_id']]
    s['current']['principal'] = s['executive']
    created_api = s['client'].post('/api/intent/dimensional/analyses/' + analysis['analysis_hash'] + '/board-packs',
                                   json={'template_id': template.template_id, 'template_version': 1,
                                         'language': 'bilingual'})
    assert created_api.status_code == 200, created_api.text
    assert created_api.json()['pack_id'] == record['pack_id']
    assert s['client'].get('/api/intent/dimensional/board-packs',
                           params={'analysis_id': analysis['analysis_hash']}).json()['packs'][0]['pack_id'] == record['pack_id']
    base = '/api/intent/dimensional/board-packs/' + record['pack_id']
    assert s['client'].get(base).json()['freshness']['status'] == 'current'
    for format in ['pdf', 'pptx']:
        response = s['client'].get(base + '/' + format)
        assert response.status_code == 200
        assert response.headers['X-Kyvern-Pack-Hash'] == record['pack_id']
        assert response.headers['X-Kyvern-Pack-Freshness'] == 'current'
    with s['connect']() as conn:
        for table in ['strategyos_intent_board_templates', 'strategyos_intent_board_packs']:
            with pytest.raises(psycopg.Error, match='immutable'):
                conn.execute(sql.SQL('UPDATE {} SET tenant_key=tenant_key').format(sql.Identifier(table)))
            conn.rollback()


def test_recorded_pack_reports_staleness_without_rewriting_snapshot(setup):
    from strategyos_mvp import board_pack, board_pack_store
    s = setup
    approve(s, import_pair(s)); analysis = analyse(s)
    template = board_pack.PackTemplate(template_id='stale-board', version=1)
    board_pack_store.register(s['operator'], template)
    record = board_pack_store.create_pack(s['executive'], analysis['analysis_hash'], 'stale-board', 1, 'en')
    original_digest = record['digest']
    newer = deepcopy(s['a']); newer['revision'] = 'newer-board-actuals'
    store.import_actuals(s['operator'], Actuals.model_validate(newer), s['pack'])
    stale = board_pack_store.read_pack(s['executive'], record['pack_id'])
    assert stale['freshness']['status'] == 'stale'
    assert stale['freshness']['reasons'] == ['new_actuals']
    assert stale['digest'] == original_digest
    assert stale['pack'] == record['pack']
    with s['connect']() as conn:
        conn.execute('UPDATE strategyos_source_access_policies SET export_allowed=false')
    with pytest.raises(PermissionError):
        board_pack_store.read_pack(s['executive'], record['pack_id'])


def test_board_template_registry_api_roles(setup):
    from strategyos_mvp import board_pack
    s = setup
    template = board_pack.PackTemplate(template_id='api-board', version=1)
    created = s['client'].post('/api/intent/dimensional/board-templates', json=template.model_dump(mode='json'))
    assert created.status_code == 200, created.text
    assert created.json()['template']['template_id'] == 'api-board'
    assert s['client'].get('/api/intent/dimensional/board-templates').json()['templates'][0]['template_id'] == 'api-board'
    s['current']['principal'] = s['executive']
    assert s['client'].post('/api/intent/dimensional/board-templates', json=template.model_dump(mode='json')).status_code == 403
    assert s['client'].get('/api/intent/dimensional/board-templates').status_code == 200
    for role in ['bu', 'system', 'analyst']:
        s['current']['principal'] = {**s['executive'], 'role': role}
        assert s['client'].get('/api/intent/dimensional/board-templates').status_code == 403


def decomposition_body(s, digest):
    source = deepcopy(s['p']['cells'][0]['source'])
    return {
        'parent_digest': digest,
        'parent_cell_id': 'regional',
        'split_dimension': 'client',
        'decimal_places': 2,
        'allocations': [
            {'cell_id': 'regional-hospital', 'member': 'hospital', 'weight': '1',
             'owner': 'hospital-owner', 'tolerance': '1', 'basis': source},
            {'cell_id': 'regional-pharmacy', 'member': 'pharmacy', 'weight': '2',
             'owner': 'pharmacy-owner', 'tolerance': '1', 'basis': source},
        ],
    }


def objective_decomposition_body(s, digest):
    return {
        'parent_digest': digest, 'metric': 'revenue',
        'split_dimensions': ['region', 'client'], 'decimal_places': 2,
        'allocations': [
            {'cell_id': 'north-hospital', 'dimensions': {'product': 'item-a', 'region': 'north', 'client': 'hospital'},
             'weight': '1', 'owner': 'North hospital lead', 'tolerance': '2', 'basis_cell_id': 'regional'},
            {'cell_id': 'north-pharmacy', 'dimensions': {'product': 'item-a', 'region': 'north', 'client': 'pharmacy'},
             'weight': '1', 'owner': 'North pharmacy lead', 'tolerance': '2', 'basis_cell_id': 'regional'},
            {'cell_id': 'south-hospital', 'dimensions': {'product': 'item-a', 'region': 'south', 'client': 'hospital'},
             'weight': '2', 'owner': 'South hospital lead', 'tolerance': '2', 'basis_cell_id': 'institutional'},
            {'cell_id': 'south-pharmacy', 'dimensions': {'product': 'item-a', 'region': 'south', 'client': 'pharmacy'},
             'weight': '4', 'owner': 'South pharmacy lead', 'tolerance': '2', 'basis_cell_id': 'institutional'},
        ],
    }


def test_whole_objective_multidimensional_decomposition_findings_and_explanation(setup):
    s = setup
    s['p']['concentration_policies'] = [{
        'policy_id': 'client-share', 'metric': 'revenue', 'dimension': 'client', 'threshold_percent': '60',
    }]
    parent = import_pair(s); approve(s, parent)
    body = objective_decomposition_body(s, parent['digest'])
    proposal = store.create_objective_decomposition(
        s['operator'], s['p']['plan_id'], 1, ObjectiveDecompositionRequest.model_validate(body))
    assert proposal['payload']['derivation']['engine_version'] == 'weighted-multidimensional-allocation.v1'
    assert proposal['payload']['derivation']['split_dimensions'] == ['region', 'client']
    assert {cell['id']: cell['target'] for cell in proposal['payload']['cells']} == {
        'north-hospital': '25.00', 'north-pharmacy': '25.00',
        'south-hospital': '50.00', 'south-pharmacy': '100.00',
    }
    prefix = f"/api/intent/dimensional/plans/{s['p']['plan_id']}/versions/1/decompose-objective"
    assert s['client'].post(prefix, json=body).json()['digest'] == proposal['digest']
    s['current']['principal'] = s['executive']
    assert s['client'].post(prefix, json=body).status_code == 403
    ratified = store.ratify(s['executive'], s['p']['plan_id'], 2, proposal['digest'],
                            'Reviewed all objective branches, owners, weights and evidence independently.')
    assert ratified['plan_digest'] == proposal['digest']
    source = deepcopy(s['a']['observations'][0]['source'])
    actuals = deepcopy(s['a'])
    actuals['revision'] = 'objective-actuals'
    actuals['observations'] = [
        {'metric': 'revenue', 'dimensions': row['dimensions'], 'unit': 'SAR', 'value': value, 'source': source}
        for row, value in zip(body['allocations'], ['0', '50', '30', '120'])
    ]
    store.import_actuals(s['operator'], Actuals.model_validate(actuals), s['pack'])
    result = store.create_analysis(s['executive'], s['p']['plan_id'], 2, 'objective-actuals', TODAY)
    offset = next(item for item in result['findings'] if item['finding_type'] == 'offset')
    assert offset['arithmetic'] == {'actual_minus_plan_behind': '-45.00',
                                    'actual_minus_plan_ahead': '45.00', 'net_variance': '0'}
    assert offset['plan_citation']['version'] == 2
    concentration = next(item for item in result['findings'] if item['finding_type'] == 'concentration')
    assert concentration['member'] == 'pharmacy'
    assert concentration['share_percent'] == '80.0'
    assert concentration['threshold_percent'] == '60'
    explanation = store.explain_cell(s['executive'], result['analysis_hash'], 'south-pharmacy')
    assert explanation['plan_citation']['version'] == 2
    assert explanation['plan_citation']['digest'] == proposal['digest']
    assert explanation['facts']['actual'] == '120'
    assert {citation['side'] for citation in explanation['evidence']} == {'plan', 'actuals'}
    api_explanation = s['client'].get('/api/intent/dimensional/analyses/' + result['analysis_hash'] +
                                      '/explain', params={'cell_id': 'south-pharmacy'})
    assert api_explanation.status_code == 200
    assert api_explanation.json()['answer'] == explanation['answer']


def test_price_volume_mix_runs_through_durable_analysis_and_board_pack(setup):
    from strategyos_mvp import board_pack
    s = setup
    source = deepcopy(s['p']['cells'][0]['source'])
    s['p']['dimensions']['product'] = ['item-a', 'item-b']
    s['p']['dimensions']['region'] = ['north']
    s['p']['dimensions']['client'] = ['retail']
    s['p']['metrics']['revenue']['planned_total'] = '300'
    s['p']['cells'] = [
        {**deepcopy(s['p']['cells'][0]), 'id': 'item-a', 'target': '100',
         'dimensions': {'product': 'item-a', 'region': 'north', 'client': 'retail'}},
        {**deepcopy(s['p']['cells'][1]), 'id': 'item-b', 'target': '200',
         'dimensions': {'product': 'item-b', 'region': 'north', 'client': 'retail'}},
    ]
    s['p']['price_volume_mix_policies'] = [{
        'bridge_id': 'commercial-bridge', 'metric': 'revenue', 'mix_dimension': 'product',
        'currency_unit': 'SAR', 'price_unit': 'SAR/unit', 'volume_unit': 'unit',
        'decimal_places': 2,
        'rows': [
            {'member': 'item-a', 'cell_id': 'item-a', 'planned_price': '10', 'planned_volume': '10',
             'price_source': source, 'volume_source': source},
            {'member': 'item-b', 'cell_id': 'item-b', 'planned_price': '20', 'planned_volume': '10',
             'price_source': source, 'volume_source': source},
        ],
    }]
    s['a']['observations'] = [
        {'metric': 'revenue', 'dimensions': {'product': 'item-a', 'region': 'north', 'client': 'retail'},
         'unit': 'SAR', 'value': '135', 'source': source},
        {'metric': 'revenue', 'dimensions': {'product': 'item-b', 'region': 'north', 'client': 'retail'},
         'unit': 'SAR', 'value': '198', 'source': source},
    ]
    s['a']['price_volume_mix'] = [{
        'bridge_id': 'commercial-bridge', 'currency_unit': 'SAR',
        'price_unit': 'SAR/unit', 'volume_unit': 'unit',
        'rows': [
            {'member': 'item-a', 'actual_price': '9', 'actual_volume': '15',
             'price_source': source, 'volume_source': source},
            {'member': 'item-b', 'actual_price': '22', 'actual_volume': '9',
             'price_source': source, 'volume_source': source},
        ],
    }]
    approve(s, import_pair(s))
    result = analyse(s)
    assert result['price_volume_mix'][0]['effects'] == {
        'volume': '60.00', 'mix': '-30.00', 'price': '3.00',
        'observed_variance': '33', 'reconstructed_variance': '33.00',
    }
    finding = next(item for item in result['findings'] if item['finding_type'] == 'price_volume_mix')
    assert finding['reconciles'] is True
    response = s['client'].get('/api/intent/dimensional/analyses/' + result['analysis_hash'])
    assert response.status_code == 200
    assert response.json()['price_volume_mix'][0]['formula_version'] == 'price-volume-mix.v1'
    pack = board_pack.compose(s['executive'], result['analysis_hash'], board_pack.PackRequest(language='bilingual'))
    assert pack['binding']['composer_version'] == 'board-pack.v3'
    assert {'plan_price', 'plan_volume', 'actual_price', 'actual_volume'} <= {
        item['side'] for item in pack['evidence']}
    bridge_page = next(page for page in pack['pages'] if 'Price / volume / mix bridge' in page['title'])
    assert any('60.00' in line for line in bridge_page['lines'])
    assert any('33.00' in line for line in bridge_page['lines'])
    assert any('[1]' in line for line in bridge_page['lines'])
    for side in ['plan_price', 'plan_volume', 'actual_price', 'actual_volume']:
        filename, content = store.evidence_bytes(s['executive'], result['analysis_hash'], 'item-a', side)
        assert filename == 'evidence.csv'
        assert content


def test_multidimensional_decomposition_rejects_unconfigured_or_undeclared_members(setup):
    s = setup
    parent = import_pair(s); approve(s, parent)
    unknown = objective_decomposition_body(s, parent['digest'])
    unknown['allocations'][0]['dimensions']['client'] = 'unknown-client'
    with pytest.raises(ValueError, match='outside the approved'):
        store.create_objective_decomposition(
            s['operator'], s['p']['plan_id'], 1, ObjectiveDecompositionRequest.model_validate(unknown))
    undeclared = objective_decomposition_body(s, parent['digest'])
    undeclared['allocations'][0]['dimensions']['product'] = 'item-b'
    with pytest.raises(ValueError, match='declaring it as a split dimension'):
        store.create_objective_decomposition(
            s['operator'], s['p']['plan_id'], 1, ObjectiveDecompositionRequest.model_validate(undeclared))
    overlapping = objective_decomposition_body(s, parent['digest'])
    overlapping['allocations'][0]['dimensions']['region'] = 'all-regions'
    overlapping['allocations'][1]['dimensions']['client'] = 'hospital'
    with pytest.raises(ValueError, match='mix parent and child'):
        store.create_objective_decomposition(
            s['operator'], s['p']['plan_id'], 1, ObjectiveDecompositionRequest.model_validate(overlapping))


def test_decomposition_requires_ratified_parent_and_is_idempotent(setup):
    s = setup
    parent = import_pair(s)
    request = DecompositionRequest.model_validate(decomposition_body(s, parent['digest']))
    with pytest.raises(store.Conflict, match='ratified'):
        store.create_decomposition(s['operator'], s['p']['plan_id'], 1, request)
    approve(s, parent)
    proposal = store.create_decomposition(s['operator'], s['p']['plan_id'], 1, request)
    repeated = store.create_decomposition(s['operator'], s['p']['plan_id'], 1, request)
    assert repeated['digest'] == proposal['digest']
    assert proposal['version'] == 2
    assert proposal['payload']['derivation']['parent_digest'] == parent['digest']
    assert proposal['payload']['derivation']['request_hash']
    children = {cell['id']: cell['target'] for cell in proposal['payload']['cells']}
    assert children['regional-hospital'] == '33.33'
    assert children['regional-pharmacy'] == '66.67'
    assert sum(Decimal(cell['target']) for cell in proposal['payload']['cells']) == Decimal('200')
    # The existing independent plan ratification makes the derived proposal authoritative.
    ratified = store.ratify(s['executive'], s['p']['plan_id'], 2, proposal['digest'],
                            'Reviewed decomposition weights, owners, evidence and exact reconciliation.')
    assert ratified['plan_digest'] == proposal['digest']
    assert store.read_plan(s['executive'], s['p']['plan_id'], 2)['governance_status'] == 'ratified'
    assert store.create_decomposition(s['operator'], s['p']['plan_id'], 1, request)['governance_status'] == 'ratified'
    split_actuals = deepcopy(s['a'])
    split_actuals['revision'] = 'decomposed-actuals'
    regional = split_actuals['observations'].pop(0)
    split_actuals['observations'].extend([
        {**regional, 'dimensions': {**regional['dimensions'], 'client': 'hospital'}, 'value': '10'},
        {**regional, 'dimensions': {**regional['dimensions'], 'client': 'pharmacy'}, 'value': '30'},
    ])
    store.import_actuals(s['operator'], Actuals.model_validate(split_actuals), s['pack'])
    result = store.create_analysis(s['executive'], s['p']['plan_id'], 2,
                                   split_actuals['revision'], TODAY)
    assert result['rollups'][0]['actual'] == '200'
    assert result['rollups'][0]['offset_detected'] is True
    from strategyos_mvp import board_pack
    pack = board_pack.compose(s['executive'], result['analysis_hash'], board_pack.PackRequest(language='en'))
    assert pack['binding']['plan_digest'] == proposal['digest']
    assert any(line.startswith('regional-hospital') for page in pack['pages'] for line in page['lines'])


def test_decomposition_api_and_generic_lineage_spoofing_boundaries(setup):
    s = setup
    parent = import_pair(s)
    approve(s, parent)
    prefix = '/api/intent/dimensional/plans/' + s['p']['plan_id'] + '/versions/1/decompose'
    body = decomposition_body(s, parent['digest'])
    response = s['client'].post(prefix, json=body)
    assert response.status_code == 200, response.text
    assert response.json()['payload']['derivation']['engine_version'] == 'weighted-allocation.v1'
    s['current']['principal'] = s['executive']
    assert s['client'].post(prefix, json=body).status_code == 403
    spoof = response.json()['payload']
    spoof['version'] = 3
    s['current']['principal'] = s['operator']
    assert s['client'].post('/api/intent/dimensional/plans',
                            json={'source_pack_id': s['pack'], 'plan': spoof}).status_code == 422


def test_decomposition_stale_parent_blocks_ratification(setup):
    s = setup
    parent = import_pair(s)
    approve(s, parent)
    request = DecompositionRequest.model_validate(decomposition_body(s, parent['digest']))
    proposal = store.create_decomposition(s['operator'], s['p']['plan_id'], 1, request)
    direct = deepcopy(s['p'])
    direct['version'] = 3
    imported = store.import_plan(s['operator'], Plan.model_validate(direct), s['pack'])
    store.ratify(s['executive'], s['p']['plan_id'], 3, imported['digest'],
                  'Reviewed a newer direct plan version and its evidence independently.')
    with pytest.raises(store.Conflict, match='newer ratified'):
        store.create_decomposition(s['operator'], s['p']['plan_id'], 1, request)
    with pytest.raises(store.Conflict, match='stale'):
        store.ratify(s['executive'], s['p']['plan_id'], 2, proposal['digest'],
                      'Reviewed the old decomposition weights, owners and evidence.')


def test_decomposition_source_and_parent_integrity_are_rechecked(setup):
    s = setup
    parent = import_pair(s)
    approve(s, parent)
    body = decomposition_body(s, parent['digest'])
    body['parent_digest'] = 'a' * 64
    with pytest.raises(store.Conflict, match='fingerprint'):
        store.create_decomposition(s['operator'], s['p']['plan_id'], 1,
                                   DecompositionRequest.model_validate(body))
    body['parent_digest'] = parent['digest']
    (s['root'] / 'raw' / 'evidence.csv').write_text('changed')
    with pytest.raises(sources.SourceUnavailable):
        store.create_decomposition(s['operator'], s['p']['plan_id'], 1,
                                   DecompositionRequest.model_validate(body))


def test_concurrent_identical_decomposition_creates_one_version(setup):
    s = setup
    parent = import_pair(s)
    approve(s, parent)
    request = DecompositionRequest.model_validate(decomposition_body(s, parent['digest']))
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(
            lambda _: store.create_decomposition(s['operator'], s['p']['plan_id'], 1, request),
            range(2),
        ))
    assert results[0]['digest'] == results[1]['digest']
    assert results[0]['version'] == results[1]['version'] == 2
    with s['connect']() as conn:
        assert conn.execute('''SELECT count(*) FROM strategyos_intent_plan_versions
            WHERE plan_id=%s''', (s['p']['plan_id'],)).fetchone()[0] == 2


def history_body(s, digest):
    return {
        'parent_digest': digest, 'parent_cell_id': 'regional', 'split_dimension': 'client',
        'historical_actual_revision': 'history-prior', 'decimal_places': 2,
        'allocations': [
            {'cell_id': 'regional-hospital', 'member': 'hospital', 'owner': 'Hospital lead',
             'tolerance': '2', 'adjustment_percent': '50'},
            {'cell_id': 'regional-pharmacy', 'member': 'pharmacy', 'owner': 'Pharmacy lead',
             'tolerance': '2', 'adjustment_percent': '0'},
        ],
    }


def import_history(s, pack=None):
    source = deepcopy(s['p']['cells'][0]['source'])
    source['locator'] = 'historical mix rows'
    year = TODAY.year - 2
    history = {
        'schema_version': 1, 'company_id': s['p']['company_id'], 'kind': 'actual',
        'revision': 'history-prior',
        'period': {'start': f'{year}-12-01', 'end': f'{year}-12-31'}, 'recorded_on': f'{year}-12-31',
        'observations': [
            {'metric': 'revenue', 'dimensions': {'product': 'item-a', 'region': 'north', 'client': 'hospital'},
             'unit': 'SAR', 'value': '40', 'source': source},
            {'metric': 'revenue', 'dimensions': {'product': 'item-a', 'region': 'north', 'client': 'pharmacy'},
             'unit': 'SAR', 'value': '60', 'source': source},
        ],
    }
    return store.import_actuals(s['operator'], Actuals.model_validate(history), pack or s['pack'])


def test_history_candidates_and_governed_decomposition_api(setup):
    s = setup
    parent = import_pair(s)
    approve(s, parent)
    history = import_history(s)
    preview = store.history_candidates(s['operator'], s['p']['plan_id'], 1, 'regional', 'client', 'history-prior')
    assert preview['readiness'] == 'ready'
    assert [row['historical_value'] for row in preview['candidates']] == ['40', '60']
    assert preview['historical_actual_digest'] == history['digest']
    request = HistoricalDecompositionRequest.model_validate(history_body(s, parent['digest']))
    proposal = store.create_history_decomposition(s['operator'], s['p']['plan_id'], 1, request)
    assert store.create_history_decomposition(s['operator'], s['p']['plan_id'], 1, request)['digest'] == proposal['digest']
    assert proposal['payload']['derivation']['historical_actual_digest'] == history['digest']
    children = {cell['id']: cell['target'] for cell in proposal['payload']['cells']}
    assert children['regional-hospital'] == children['regional-pharmacy'] == '50.00'
    prefix = f"/api/intent/dimensional/plans/{s['p']['plan_id']}/versions/1"
    params = {'parent_cell_id': 'regional', 'split_dimension': 'client', 'actual_revision': 'history-prior'}
    assert s['client'].get(prefix + '/history-candidates', params=params).json()['readiness'] == 'ready'
    assert s['client'].post(prefix + '/decompose-from-history', json=history_body(s, parent['digest'])).status_code == 200
    s['current']['principal'] = s['executive']
    assert s['client'].get(prefix + '/history-candidates', params=params).status_code == 403
    assert s['client'].post(prefix + '/decompose-from-history', json=history_body(s, parent['digest'])).status_code == 403
    ratified = store.ratify(s['executive'], s['p']['plan_id'], 2, proposal['digest'],
                            'Reviewed the historical mix, adjustments, owners and evidence lineage.')
    assert ratified['plan_digest'] == proposal['digest']


def test_history_decomposition_blocks_missing_history_and_stale_parent(setup):
    s = setup
    parent = import_pair(s)
    approve(s, parent)
    import_history(s)
    body = history_body(s, parent['digest'])
    body['allocations'][1]['member'] = 'missing-member'
    with pytest.raises(ValueError, match='No historical observation'):
        store.create_history_decomposition(s['operator'], s['p']['plan_id'], 1,
                                           HistoricalDecompositionRequest.model_validate(body))
    direct = deepcopy(s['p'])
    direct['version'] = 2
    newer = store.import_plan(s['operator'], Plan.model_validate(direct), s['pack'])
    store.ratify(s['executive'], s['p']['plan_id'], 2, newer['digest'],
                  'Reviewed this newer direct plan and its complete evidence set.')
    with pytest.raises(store.Conflict, match='newer ratified'):
        store.create_history_decomposition(s['operator'], s['p']['plan_id'], 1,
                                           HistoricalDecompositionRequest.model_validate(history_body(s, parent['digest'])))


def advisor_body(s, digest):
    return {
        'schema_version': 1, 'config_id': 'synthetic-client-setup', 'version': 1,
        'executive_sponsor': 'Chief Executive',
        'objective': 'Expose client-level revenue drift against the plan approved by the board.',
        'plan_id': s['p']['plan_id'], 'plan_version': 1, 'plan_digest': digest,
        'parent_cell_id': 'regional', 'split_dimension': 'client',
        'historical_actual_revision': 'history-prior', 'decimal_places': 2,
        'client': {'en': 'Synthetic Healthcare', 'ar': 'الرعاية الصحية التجريبية'},
        'board_title': {'en': 'Revenue performance review', 'ar': 'مراجعة أداء الإيرادات'},
        'metric_label': {'en': 'Revenue', 'ar': 'الإيرادات'},
        'dimension_label': {'en': 'Client', 'ar': 'العميل'},
        'additional_labels': {
            'product': {'en': 'Product', 'ar': 'المنتج'}, 'region': {'en': 'Region', 'ar': 'المنطقة'},
            'item-a': {'en': 'Item A', 'ar': 'الصنف أ'}, 'north': {'en': 'North', 'ar': 'الشمال'},
            'south': {'en': 'South', 'ar': 'الجنوب'}, 'institution': {'en': 'Institution', 'ar': 'مؤسسة'},
        },
        'allocations': [
            {'cell_id': 'regional-hospital', 'member': 'hospital', 'owner': 'Hospital lead',
             'tolerance': '2', 'adjustment_percent': '50',
             'label': {'en': 'Hospital', 'ar': 'مستشفى'}},
            {'cell_id': 'regional-pharmacy', 'member': 'pharmacy', 'owner': 'Pharmacy lead',
             'tolerance': '2', 'adjustment_percent': '0',
             'label': {'en': 'Pharmacy', 'ar': 'صيدلية'}},
        ],
    }


def test_advisor_configuration_is_versioned_approved_and_published_without_source_override(setup):
    s = setup
    parent = import_pair(s)
    approve(s, parent)
    history_pack = 'owned-history-pack'
    history_root = s['root'].parent / history_pack
    shutil.copytree(s['root'], history_root)
    history_manifest = json.loads((history_root / 'summary.json').read_text())
    history_manifest['source_pack_id'] = history_pack
    (history_root / 'summary.json').write_text(json.dumps(history_manifest))
    import_history(s, history_pack)
    body = advisor_body(s, parent['digest'])
    configured = advisor_config_store.create(s['operator'], AdvisorConfiguration.model_validate(body))
    assert configured['readiness']['status'] == 'ready'
    assert all(configured['readiness']['checks'].values())
    assert configured['source_bindings']['plan']['source_pack_id'] == s['pack']
    assert configured['source_bindings']['historical_actuals']['source_pack_id'] == history_pack
    assert configured['approval'] is None
    with pytest.raises(PermissionError):
        advisor_config_store.approve(s['operator'], body['config_id'], 1, configured['digest'],
                                     'Reviewed the complete client configuration and mappings.')
    approval = advisor_config_store.approve(s['executive'], body['config_id'], 1, configured['digest'],
                                            'Reviewed the complete client configuration and bilingual mappings.')
    assert approval['config_digest'] == configured['digest']
    template = advisor_config_store.template(s['executive'], body['config_id'], 1)
    assert template['client']['ar'] == body['client']['ar']
    assert template['labels']['hospital']['en'] == 'Hospital'
    from strategyos_mvp import board_pack_store
    registered_template = board_pack_store.register_advisor(s['operator'], body['config_id'], 1)
    assert registered_template['origin'] == {'type': 'advisor', 'config_id': body['config_id'],
                                              'version': 1, 'config_digest': configured['digest']}
    assert registered_template['template'] == template
    publication = advisor_config_store.publish(s['operator'], body['config_id'], 1, configured['digest'])
    assert publication['plan_version'] == 2
    assert advisor_config_store.publish(s['operator'], body['config_id'], 1, configured['digest']) == publication
    proposal = store.read_plan(s['executive'], s['p']['plan_id'], 2)
    assert proposal['payload']['derivation']['engine_version'] == 'history-adjusted-allocation.v1'
    assert advisor_config_store.catalog(s['executive'])['configurations'][0]['approved'] is True
    assert advisor_config_store.catalog(s['executive'])['configurations'][0]['published'] is True
    with s['connect']() as conn:
        for table in ['strategyos_intent_advisor_configs', 'strategyos_intent_advisor_approvals',
                      'strategyos_intent_advisor_publications']:
            with pytest.raises(psycopg.Error, match='immutable'):
                conn.execute(sql.SQL('UPDATE {} SET tenant_key=tenant_key').format(sql.Identifier(table)))
            conn.rollback()


def test_advisor_configuration_api_roles_and_complete_mapping(setup):
    s = setup
    parent = import_pair(s)
    approve(s, parent)
    import_history(s)
    body = advisor_body(s, parent['digest'])
    body['allocations'][1]['member'] = 'member-without-history'
    response = s['client'].post('/api/intent/dimensional/advisor/configurations', json=body)
    assert response.status_code == 422
    assert 'incomplete' in response.json()['detail']
    body = advisor_body(s, parent['digest'])
    spoof = deepcopy(body)
    spoof['source_pack_id'] = 'caller-selected-pack'
    assert s['client'].post('/api/intent/dimensional/advisor/configurations', json=spoof).status_code == 422
    incomplete_labels = deepcopy(body)
    incomplete_labels['additional_labels'].pop('region')
    missing_label = s['client'].post('/api/intent/dimensional/advisor/configurations', json=incomplete_labels)
    assert missing_label.status_code == 422
    assert 'Bilingual board labels are incomplete' in missing_label.json()['detail']
    configured = s['client'].post('/api/intent/dimensional/advisor/configurations', json=body)
    assert configured.status_code == 200, configured.text
    record = configured.json()
    s['current']['principal'] = s['executive']
    prefix = '/api/intent/dimensional/advisor/configurations/synthetic-client-setup/versions/1'
    assert s['client'].post('/api/intent/dimensional/advisor/configurations', json=body).status_code == 403
    approved = s['client'].post(prefix + '/approve', json={
        'expected_digest': record['digest'], 'note': 'Reviewed all guided fields and source-bound mappings independently.'})
    assert approved.status_code == 200, approved.text
    assert s['client'].get(prefix + '/board-template').status_code == 200
    assert s['client'].post(prefix + '/board-template/register', json={}).status_code == 403
    assert s['client'].post(prefix + '/publish', json={'expected_digest': record['digest']}).status_code == 403
    s['current']['principal'] = s['operator']
    registered = s['client'].post(prefix + '/board-template/register', json={})
    assert registered.status_code == 200, registered.text
    assert registered.json()['origin']['type'] == 'advisor'
    assert s['client'].post(prefix + '/publish', json={'expected_digest': record['digest']}).status_code == 200
