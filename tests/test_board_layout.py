from io import BytesIO
import json

import pytest
from pypdf import PdfReader
from pptx import Presentation

from strategyos_mvp import board_pack
from strategyos_mvp.board_layout import table_pages, wrap, width


@pytest.mark.parametrize('language', ['en', 'ar', 'bilingual'])
def test_full_year_table_retains_all_576_rows_and_exact_evidence_in_both_exports(language):
    headers = [board_pack.word(key, language) for key in ['cells', 'owner', 'target', 'actual', 'variance', 'status', 'evidence']]
    rows = [[f'Revenue (SAR)\n2026-01 · Central · Hospital/private · Product {i}',
             'Lina Al-Qahtani — Central Region GM (Sales)', f'{10000000+i}.01', f'{10000000+i}.02',
             '0.01', board_pack.word('ahead', language), f'[{i+1}]'] for i in range(576)]
    links = [{'2': f'/evidence/{i}/plan', '3': f'/evidence/{i}/actual'} for i in range(576)]
    pages = table_pages(board_pack.word('cells', language), headers, [250,140,100,100,100,90,100], rows, links=links)
    assert [row for page in pages for row in page['table']['rows']] == rows
    assert all(sum(p['table']['row_heights']) + p['table']['header_height'] <= 372 for p in pages)
    evidence = [{'number': i+1, 'path': links[i]['2']} for i in range(576)]
    pack = {'pack_hash': 'synthetic-pack', 'binding': {'template': board_pack.PackTemplate().model_dump(),
            'language': language}, 'checked_at': '2026-09-12', 'pages': pages, 'evidence': evidence}
    pdf = PdfReader(BytesIO(board_pack.export_pdf(pack, 'https://example.test')))
    deck = Presentation(BytesIO(board_pack.export_pptx(pack, 'https://example.test')))
    assert len(pdf.pages) == len(deck.slides) == len(pages)
    assert json.loads(pdf.attachments['kyvern-evidence-register.json'][0])['evidence'] == evidence
    assert json.loads(deck.slides[0].notes_slide.notes_text_frame.text)['evidence'] == evidence
    exported = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
    # Every exact value remains editable, once and in order; no picture flattening.
    actuals = [table.cell(row,3).text.replace('\u200e','') for table in exported for row in range(1,len(table.rows))]
    assert actuals == [row[3] for row in rows]
    assert sum(len(page.get('/Annots', [])) for page in pdf.pages) == 1152


def test_long_words_wrap_without_truncation():
    value = 'source_' * 120
    lines = wrap(value, 80, 11)
    assert ''.join(lines) == value
    assert all(width(line, 11) <= 80 for line in lines)
