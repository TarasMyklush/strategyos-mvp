"""Deterministic board packs over immutable Intent analyses; no model-authored numbers.

Templates are portable client pack artifacts. Exports are snapshots, not live decks.
Source export authorization and byte integrity are rechecked on every composition.
"""
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Literal
from urllib.parse import urlencode, quote
import re

from pydantic import Field, field_validator

from . import dimensional_intent_store as store
from .dimensional_plan import (
    Actuals, Contract, Plan, actual_source_references, fingerprint,
    plan_source_references,
)
from .dimensional_intent_sources import registered_sources


class Translation(Contract):
    en: str = Field(min_length=1, max_length=100)
    ar: str = Field(min_length=1, max_length=100)

    @field_validator('en', 'ar')
    @classmethod
    def printable(cls, value):
        if any(ord(c) < 32 or c in '\u202a\u202b\u202c\u202d\u202e\u2066\u2067\u2068\u2069' for c in value):
            raise ValueError('Template text must be a single line without direction overrides.')
        return value.strip()


class PackTemplate(Contract):
    schema_version: Literal[1] = 1
    template_id: str = Field(default='kyvern-board', pattern=r'^[A-Za-z0-9][A-Za-z0-9_.-]{0,79}$')
    version: int = Field(default=1, ge=1, strict=True)
    title: Translation = Field(default_factory=lambda: Translation(en='Board performance review', ar='مراجعة الأداء لمجلس الإدارة'))
    client: Translation = Field(default_factory=lambda: Translation(en='Kyvern', ar='كايفرن'))
    accent: str = Field(default='#72533F', pattern=r'^#[0-9a-fA-F]{6}$')
    labels: dict[str, Translation] = Field(default_factory=dict, max_length=200)


class PackRequest(Contract):
    template: PackTemplate = Field(default_factory=PackTemplate)
    language: Literal['en', 'ar', 'bilingual'] = 'bilingual'


WORDS = {
    'summary': ('Performance summary', 'ملخص الأداء'), 'cells': ('Plan cells', 'بنود الخطة'),
    'evidence': ('Evidence references', 'مراجع الأدلة'), 'target': ('Target', 'المستهدف'),
    'actual': ('Actual', 'الفعلي'), 'variance': ('Variance', 'الانحراف'),
    'missing': ('Missing', 'غير متوفر'), 'on_plan': ('On plan', 'وفق الخطة'),
    'ahead': ('Ahead', 'متقدم'), 'behind': ('Behind', 'متأخر'), 'incomplete': ('Incomplete', 'غير مكتمل'),
    'owner': ('Owner', 'المسؤول'), 'status': ('Status', 'الحالة'),
    'offset': ('An aggregate result masks cells behind plan.', 'النتيجة الإجمالية تخفي بنوداً متأخرة عن الخطة.'),
    'coverage': ('Measured / planned cells', 'البنود المقاسة / المخططة'),
    'new_plan': ('A newer ratified plan exists for this period. Re-run the comparison to use it.', 'توجد خطة معتمدة أحدث لهذه الفترة. أعد المقارنة لاستخدامها.'),
    'new_actuals': ('Newer actual snapshots exist for this period. Review their scope before selecting a replacement.', 'توجد لقطات فعلية أحدث لهذه الفترة. راجع نطاقها قبل اختيار بديل.'),
    'snapshot': ('Fixed snapshot. Downloaded files do not refresh automatically.', 'لقطة ثابتة. الملفات المحملة لا تتحدث تلقائياً.'),
    'assurance': ('Evidence file integrity checked; source values are not independently certified.', 'تم التحقق من سلامة ملفات الأدلة؛ لم يتم التصديق على قيم المصادر بشكل مستقل.'),
    'unplanned': ('Unplanned actual tuples: ', 'بنود فعلية خارج الخطة: '),
    'finding': ('Evidence-bound finding', 'نتيجة مرتبطة بالأدلة'),
    'price_volume_mix': ('Price / volume / mix bridge', 'تحليل السعر / الحجم / المزيج'),
    'plan_revenue': ('Planned revenue', 'الإيراد المخطط'),
    'actual_revenue': ('Actual revenue', 'الإيراد الفعلي'),
    'volume_effect': ('Volume effect', 'أثر الحجم'),
    'mix_effect': ('Mix effect', 'أثر المزيج'),
    'price_effect': ('Price effect', 'أثر السعر'),
    'observed_variance': ('Observed variance', 'الانحراف المرصود'),
    'reconstructed_variance': ('Reconstructed variance', 'الانحراف المعاد بناؤه'),
    'reconciliation': ('Exact reconciliation', 'مطابقة حسابية تامة'),
    'concentration': ('Concentration above threshold', 'التركيز أعلى من الحد'),
    'share': ('Share', 'الحصة'),
    'threshold': ('Threshold', 'الحد'),
    'allocation_basis': ('Approved allocation basis', 'أساس التوزيع المعتمد'),
    'history': ('Historical value', 'القيمة التاريخية'),
    'seasonality': ('Seasonal factor', 'المعامل الموسمي'),
    'adjustment': ('Planning adjustment', 'تعديل التخطيط'),
    'effective_weight': ('Effective weight', 'الوزن الفعلي'),
    'not_applied': ('Not applied', 'غير مطبق'),
}


def word(key, language):
    pair = WORDS[key]
    return pair[0] if language == 'en' else pair[1] if language == 'ar' else pair[0] + ' | ' + pair[1]


def translated(value, language):
    return value.en if language == 'en' else value.ar if language == 'ar' else value.en + ' | ' + value.ar


def reference_numbers(numbers):
    """Lossless compact ranges; never truncate a finding's contributing sources."""
    ranges = []
    for n in sorted(set(numbers)):
        if ranges and n == ranges[-1][1] + 1:
            ranges[-1][1] = n
        else:
            ranges.append([n, n])
    return ' '.join('[' + (str(a) if a == b else f'{a}–{b}') + ']' for a, b in ranges)


def compose(principal, analysis_id, request):
    tenant, _ = store._scope(principal)
    result = store.read_analysis(principal, analysis_id)
    # Bound work and export size. Never silently truncate a board report.
    if len(result['cells']) > 10000 or len(result['rollups']) > 200:
        raise ValueError('This reporting range exceeds the export budget of 10,000 cells or 200 metrics. Select a smaller reporting range.')
    with store._connection() as conn:
        plan = store._plan(conn, tenant, result['plan_id'], result['plan_version'])
        actual = store._actuals(conn, tenant, result['actual_revision'])
        cursor = conn.execute('''SELECT p.* FROM strategyos_intent_plan_versions p
            JOIN strategyos_intent_ratifications r USING(tenant_key,plan_id,version)
            WHERE p.tenant_key=%s AND p.plan_id=%s AND p.version>%s
            AND p.payload->'period'=%s::jsonb ORDER BY p.version DESC''',
            (tenant, plan['plan_id'], plan['version'], store._encode(plan['payload']['period'])))
        newer_plans = [dict(zip([c.name for c in cursor.description], r)) for r in cursor.fetchall()]
        cursor = conn.execute('''SELECT * FROM strategyos_intent_actual_versions
            WHERE tenant_key=%s AND imported_at>%s AND payload->'period'=%s::jsonb ORDER BY imported_at DESC''',
            (tenant, actual['imported_at'], store._encode(actual['payload']['period'])))
        newer_actuals = [dict(zip([c.name for c in cursor.description], r)) for r in cursor.fetchall()]
    for row, kind in [(plan, 'plan'), (actual, 'actuals')]:
        store._sources(principal, row, kind=kind, purpose='export')
    # Do not disclose revisions from a source whose access has been revoked.
    from .dimensional_intent_sources import SourceUnavailable
    warnings, newer_records = [], {}
    for key, candidates, kind in [('new_plan', newer_plans, 'plan'), ('new_actuals', newer_actuals, 'actuals')]:
        for row in candidates:
            store._checked(row)
            try:
                store._sources(principal, row, kind=kind, verify_bytes=False)
            except (PermissionError, SourceUnavailable):
                continue
            warnings.append(key)
            newer_records[key] = ({'plan_id': row['plan_id'], 'version': row['version'], 'digest': row['digest']}
                                  if kind == 'plan' else {'revision': row['revision'], 'digest': row['digest']})
            break
    return compose_snapshot(
        result, request, analysis_id=analysis_id,
        plan_digest=result['plan_import_digest'], actual_digest=result['actual_import_digest'],
        warnings=warnings, newer_records=newer_records,
        plan_derivation=plan['payload'].get('derivation'),
    )


def compose_snapshot(result, request, *, analysis_id, plan_digest, actual_digest,
                     warnings=None, newer_records=None, evidence_url=None, plan_derivation=None):
    """Compose the same deterministic pages from an already verified analysis snapshot."""
    if len(result['cells']) > 10000 or len(result['rollups']) > 200:
        raise ValueError('This reporting range exceeds the export budget of 10,000 cells or 200 metrics. Select a smaller reporting range.')
    warnings = list(warnings or [])
    newer_records = dict(newer_records or {})
    if evidence_url is None:
        evidence_url = lambda side, cell_id: (
            '/api/intent/dimensional/analyses/' + analysis_id + '/evidence?' +
            urlencode({'side': side, 'cell_id': cell_id}))
    lang, template = request.language, request.template
    def label(value):
        return translated(template.labels[value], lang) if value in template.labels else str(value).replace('_', ' ')
    def amount(value):
        return word('missing', lang) if value is None else value
    pages = []
    def add(title, lines, links=None):
        from .board_layout import wrap
        expanded = [(part, (links or {}).get(str(i))) for i, line in enumerate(lines)
                    for part in wrap(line, 880, 17)]
        for start in range(0, len(expanded), 9):
            chunk = expanded[start:start + 9]
            pages.append({'title': title, 'lines': [text for text, _ in chunk],
                          'links': {str(i): link for i, (_, link) in enumerate(chunk) if link}})
    add(translated(template.title, lang), [translated(template.client, lang),
        f"Plan version / إصدار الخطة: {result['plan_version']}",
        f"{result['period']['start']} — {result['period']['end']}",
        f"As of / كما في: {result['as_of']}",
        word('snapshot', lang), word('assurance', lang)])
    for key in warnings:
        add(word('summary', lang), [word(key, lang)])
    for rollup in result['rollups']:
        lines = [label(rollup['metric']) + ' (' + rollup['unit'] + ')',
            word('target', lang) + ': ' + amount(rollup['target']),
            word('actual', lang) + ': ' + amount(rollup['actual']),
            word('variance', lang) + ': ' + amount(rollup['variance']),
            word('status', lang) + ': ' + word(rollup['status'], lang),
            word('coverage', lang) + ': ' + str(rollup['measured_cells']) + ' / ' + str(rollup['planned_cells'])]
        if rollup['offset_detected']: lines.append(word('offset', lang))
        if rollup['unplanned_actuals']: lines.append(word('unplanned', lang) + str(len(rollup['unplanned_actuals'])))
        add(word('summary', lang), lines)
    evidence = []
    cell_rows, cell_keys, cell_links = [], [], []
    for c in result['cells']:
        refs = []
        for side, key in [('plan', 'plan_source'), ('actuals', 'actual_source')]:
            source = c[key]
            if source:
                number = len(evidence) + 1
                path = evidence_url(side, c['cell_id'])
                evidence.append({'number': number, 'cell_id': c['cell_id'], 'side': side, 'source': source, 'path': path})
                refs.append('[' + str(number) + ']')
        cell_keys.append(c['cell_id'])
        cell_rows.append([
            label(c['metric']) + ' (' + c['unit'] + ')\n' + ' · '.join(label(v) for k, v in sorted(c['dimensions'].items())),
            c['owner'], amount(c['target']), amount(c['actual']), amount(c['variance']),
            word(c['status'], lang), ' '.join(refs),
        ])
        cell_links.append({str(column): evidence_url(side, c['cell_id'])
                           for column, side, source in [(2, 'plan', c['plan_source']), (3, 'actuals', c['actual_source'])]
                           if source})
    for bridge in result.get('price_volume_mix', []):
        if bridge.get('status') != 'reconciled':
            continue
        for row in bridge['rows']:
            for side, key in [
                ('plan_price', 'plan_price_source'), ('plan_volume', 'plan_volume_source'),
                ('actual_price', 'actual_price_source'), ('actual_volume', 'actual_volume_source'),
            ]:
                source = row[key]
                evidence.append({
                    'number': len(evidence) + 1, 'cell_id': row['cell_id'], 'side': side,
                    'source': source, 'path': evidence_url(side, row['cell_id']),
                })
    evidence_numbers = {(item['cell_id'], item['side']): item['number'] for item in evidence}
    for finding in result.get('findings', []):
        finding_type = finding.get('finding_type')
        refs = []
        for cell in finding.get('cells', []):
            for side in ('plan', 'actuals'):
                number = evidence_numbers.get((cell.get('cell_id'), side))
                if number is not None and number not in refs:
                    refs.append(number)
        if finding_type == 'price_volume_mix':
            effects = finding['effects']
            refs = [number for cell in finding['cells'] for side in (
                'plan', 'actuals', 'plan_price', 'plan_volume', 'actual_price', 'actual_volume')
                if (number := evidence_numbers.get((cell['cell_id'], side))) is not None]
            lines = [
                label(finding['metric']) + ' (' + finding['currency_unit'] + ')',
                word('plan_revenue', lang) + ': ' + finding['plan']['revenue'],
                word('actual_revenue', lang) + ': ' + finding['actual']['revenue'],
                word('volume_effect', lang) + ': ' + effects['volume'],
                word('mix_effect', lang) + ': ' + effects['mix'],
                word('price_effect', lang) + ': ' + effects['price'],
                word('observed_variance', lang) + ': ' + effects['observed_variance'],
                word('reconstructed_variance', lang) + ': ' + effects['reconstructed_variance'],
                word('reconciliation', lang) + ': ' + ('✓' if finding['reconciles'] else '✕'),
                word('evidence', lang) + ': ' + reference_numbers(refs),
            ]
            add(word('price_volume_mix', lang), lines)
        elif finding_type == 'offset':
            arithmetic = finding['arithmetic']
            add(word('finding', lang), [
                label(finding['metric']), word('offset', lang),
                word('behind', lang) + ': ' + arithmetic['actual_minus_plan_behind'],
                word('ahead', lang) + ': ' + arithmetic['actual_minus_plan_ahead'],
                word('variance', lang) + ': ' + arithmetic['net_variance'],
                word('evidence', lang) + ': ' + reference_numbers(refs),
            ])
        elif finding_type == 'concentration':
            refs = [number for side in ('plan', 'actuals')
                    if (number := evidence_numbers.get((finding['cell_id'], side))) is not None]
            add(word('concentration', lang), [
                label(finding['metric']) + ' · ' + label(finding['dimension']) + ': ' + label(finding['member']),
                word('share', lang) + ': ' + finding['share_percent'] + '%',
                word('threshold', lang) + ': ' + finding['threshold_percent'] + '%',
                word('evidence', lang) + ': ' + reference_numbers(refs),
            ])
    from .board_layout import table_pages
    pages.extend(table_pages(word('cells', lang),
        [word('cells', lang), word('owner', lang), word('target', lang), word('actual', lang),
         word('variance', lang), word('status', lang), word('evidence', lang)],
        [250, 140, 100, 100, 100, 90, 100], cell_rows, keys=cell_keys, links=cell_links))
    if plan_derivation and plan_derivation.get('historical_actual_revision'):
        by_cell = {cell['cell_id']: cell for cell in result['cells']}
        allocation_rows, allocation_links, allocation_keys = [], [], []
        for allocation in plan_derivation['allocations']:
            cell = by_cell.get(allocation['cell_id'])
            if cell is None:
                continue
            links, refs = {}, []
            for column, side, source in [(1,'history',allocation['basis']),
                                         (2,'seasonality',allocation.get('seasonality_basis'))]:
                if not source:
                    continue
                path = '/api/intent/dimensional/plans/' + quote(result['plan_id'],safe='') + '/versions/' + str(result['plan_version']) + '/evidence?' + urlencode({'cell_id':cell['cell_id'],'basis':side})
                number = len(evidence) + 1
                evidence.append({'number':number,'cell_id':cell['cell_id'],'side':side,'source':source,'path':path})
                refs.append(number); links[str(column)] = path
            allocation_rows.append([label(cell['metric']) + '\n' + ' · '.join(label(v) for k,v in sorted(cell['dimensions'].items())),
                allocation['historical_value'] + ' ' + cell['unit'], allocation.get('seasonality_factor') or word('not_applied',lang),
                allocation['adjustment_percent'] + '%', allocation['effective_weight'] + ' ' + cell['unit'],
                cell['target'] + ' ' + cell['unit'], reference_numbers(refs)])
            allocation_links.append(links); allocation_keys.append(cell['cell_id'])
        pages.extend(table_pages(word('allocation_basis',lang),
            [word('cells',lang), word('history',lang), word('seasonality',lang), word('adjustment',lang),
             word('effective_weight',lang), word('target',lang), word('evidence',lang)],
            [250,100,90,100,110,110,120], allocation_rows,keys=allocation_keys,links=allocation_links))
    # Each value links to its exact governed evidence. Preserve the full register
    # in the PDF attachment/PPTX notes instead of hundreds of one-source slides.
    documents = {}
    for e in evidence:
        key = (e['source']['path'], e['source']['sha256'])
        documents.setdefault(key, []).append(e)
    for number, entries in enumerate(documents.values(), 1):
        add(word('evidence', lang) + ' · ' + str(number), [
            word('evidence', lang) + ': ' + str(len(entries)),
            'Open verified source / افتح المصدر الموثق',
            'Exact row references are retained in the attached evidence register. / مراجع الصفوف محفوظة في سجل الأدلة المرفق.',
        ], {'1': entries[0]['path']})
    binding = {'composer_version': 'board-pack.v4', 'analysis_hash': analysis_id, 'plan_digest': plan_digest,
               'actual_digest': actual_digest, 'template': template.model_dump(mode='json'),
               'language': lang, 'warnings': warnings, 'newer_records': newer_records}
    if plan_derivation:
        binding['plan_derivation'] = plan_derivation
    add('Snapshot references / مراجع اللقطة', [
        word('snapshot', lang), word('assurance', lang),
        'Technical provenance is retained in the PDF attachment and slide notes.',
        'حُفظت المراجع التقنية في مرفق ملف PDF وملاحظات الشرائح.',
    ])
    return {'schema_version': 1, 'pack_hash': fingerprint(binding), 'binding': binding,
            'checked_at': datetime.now(timezone.utc).isoformat(), 'pages': pages, 'evidence': evidence,
            'untranslated_labels': sorted({v for c in result['cells'] for v in [c['metric'], *c['dimensions'].keys(), *c['dimensions'].values()] if v not in template.labels}) if lang != 'en' else []}


FONT_ROOT = Path(__file__).with_name('fonts')


def mark_ltr(text):
    return re.sub(r'[+-]?[A-Za-z0-9][A-Za-z0-9_.:/-]*', lambda m: '\u200e' + m.group() + '\u200e', text)


def visual(text):
    import arabic_reshaper
    from bidi.algorithm import get_display
    # Keep dates, signs and exact decimal strings in logical LTR order inside RTL prose.
    text = mark_ltr(text)
    return get_display(arabic_reshaper.reshape(text), base_dir='L').replace('\u200e', '')


def export_pdf(pack, public_url):
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.colors import HexColor
    for name, file in [('BoardLatin', 'NotoSans-Regular.ttf'), ('BoardArabic', 'NotoSansArabic-Regular.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont(name, str(FONT_ROOT / file)))
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=(960, 540))
    pdf.setTitle(pack['binding']['template']['title']['en'])
    def draw(text, x, y, size, width):
        # Arabic font covers Arabic and Latin; preserve numeric strings without conversion.
        text = visual(text)
        arabic = pdfmetrics.getFont('BoardArabic').face.charToGlyph
        runs = []
        for char in text:
            font = 'BoardArabic' if ord(char) in arabic and ord(char) > 255 else 'BoardLatin'
            if runs and runs[-1][0] == font: runs[-1][1] += char
            else: runs.append([font, char])
        def length(): return sum(pdfmetrics.stringWidth(value, font, size) for font, value in runs)
        while length() > width and size > 8: size -= .5
        if length() > width:
            raise ValueError('Report text is too long for export. Shorten the template or plan labels.')
        for font, value in runs:
            pdf.setFont(font, size)
            pdf.drawString(x, y, value)
            x += pdfmetrics.stringWidth(value, font, size)
    for i, page in enumerate(pack['pages']):
        pdf.setFillColor(HexColor('#FBF8F2')); pdf.rect(0, 0, 960, 540, fill=1, stroke=0)
        pdf.setFillColor(HexColor(pack['binding']['template']['accent']))
        draw(page['title'], 40, 478, 28, 880)
        pdf.setFillColor(HexColor('#192C35'))
        if page.get('table'):
            from .board_layout import wrap
            table = page['table']
            y = 440
            rows = [table['headers'], *table['rows']]
            heights = [table['header_height'], *table['row_heights']]
            for row_index, (row, height) in enumerate(zip(rows, heights)):
                x = 40
                for column, (value, column_width) in enumerate(zip(row, table['widths'])):
                    for line_index, line in enumerate(wrap(value, column_width - 12, table['font_size'])):
                        draw(line, x + 6, y - 17 - line_index * 15, table['font_size'], column_width - 12)
                    link = table['links'][row_index - 1].get(str(column)) if row_index else None
                    if link and public_url:
                        pdf.linkURL(public_url + link, (x, y-height, x+column_width, y), relative=0)
                    x += column_width
                y -= height
                pdf.setStrokeColor(HexColor('#D8D4CC')); pdf.setLineWidth(.5)
                pdf.line(40, y, 920, y)
        else:
            for n, line in enumerate(page['lines']):
                y = 426 - n * 43
                draw(line, 40, y, 17, 880)
                if str(n) in page['links'] and public_url:
                    pdf.linkURL(public_url + page['links'][str(n)], (40, y-5, 920, y+22), relative=0)
        draw(f"Kyvern · {i+1}/{len(pack['pages'])}", 40, 28, 8, 880)
        pdf.showPage()
    pdf.save()
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    writer.clone_document_from_reader(PdfReader(BytesIO(output.getvalue())))
    writer.add_attachment('kyvern-evidence-register.json', store._encode({
        'pack_hash': pack['pack_hash'], 'binding': pack['binding'], 'evidence': pack['evidence'],
    }).encode('utf-8'))
    attached = BytesIO(); writer.write(attached)
    return attached.getvalue()


def export_pptx(pack, public_url):
    # Runtime export lives in the Python API; editable text retains Unicode/RTL.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.oxml.xmlchemy import OxmlElement
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    def text(slide, value, y, size, color, link=None):
        shape = slide.shapes.add_textbox(Inches(.55), Inches(y), Inches(12.2), Inches(.55))
        tf = shape.text_frame; tf.word_wrap = False; tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        p = tf.paragraphs[0]
        run = p.add_run(); run.text = mark_ltr(value)
        run.font.name = 'Noto Sans'
        run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color)
        if pack['binding']['language'] == 'ar': p._p.get_or_add_pPr().set('rtl', '1')
        if re.search(r'[\u0600-\u06ff]', value):
            cs = OxmlElement('a:cs'); cs.set('typeface', 'Noto Sans Arabic'); run._r.get_or_add_rPr().append(cs)
        if link and public_url: run.hyperlink.address = public_url + link
    for i, page in enumerate(pack['pages']):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string('FBF8F2')
        text(slide, page['title'], .4, 28, pack['binding']['template']['accent'][1:])
        if page.get('table'):
            from .board_layout import wrap
            data = page['table']
            heights = [data['header_height'], *data['row_heights']]
            shape = slide.shapes.add_table(len(data['rows']) + 1, len(data['headers']),
                Pt(40), Pt(100), Pt(880), Pt(sum(heights)))
            table = shape.table
            for column, width in zip(table.columns, data['widths']): column.width = Pt(width)
            for row_index, (values, height) in enumerate(zip([data['headers'], *data['rows']], heights)):
                table.rows[row_index].height = Pt(height)
                for column, (value, width) in enumerate(zip(values, data['widths'])):
                    cell = table.cell(row_index, column)
                    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string('FBF8F2')
                    cell.margin_left = cell.margin_right = Pt(6)
                    cell.margin_top = cell.margin_bottom = Pt(6)
                    tf = cell.text_frame; tf.clear(); tf.word_wrap = False
                    for line_index, line in enumerate(wrap(value, width - 12, data['font_size'])):
                        p = tf.paragraphs[0] if line_index == 0 else tf.add_paragraph()
                        p.space_before = p.space_after = Pt(0); p.line_spacing = Pt(15)
                        if pack['binding']['language'] == 'ar': p._p.get_or_add_pPr().set('rtl', '1')
                        run = p.add_run(); run.text = mark_ltr(line)
                        run.font.name = 'Noto Sans'; run.font.size = Pt(data['font_size'])
                        run.font.bold = False; run.font.color.rgb = RGBColor.from_string('192C35')
                        cs = OxmlElement('a:cs'); cs.set('typeface', 'Noto Sans Arabic'); run._r.get_or_add_rPr().append(cs)
                        link = data['links'][row_index - 1].get(str(column)) if row_index else None
                        if link and public_url: run.hyperlink.address = public_url + link
        else:
            for n, line in enumerate(page['lines']):
                size = min(17, max(8, 1000 / max(len(line), 1)))
                text(slide, line, 1.22 + n * .597, size, '192C35', page['links'].get(str(n)))
        text(slide, f"Kyvern · {i+1}/{len(pack['pages'])}", 6.95, 8, '72533F')
        provenance = {'pack_hash': pack['pack_hash'], 'binding': pack['binding'], 'checked_at': pack['checked_at']}
        if i == 0: provenance['evidence'] = pack['evidence']
        if page.get('table'): provenance['cell_ids'] = page['table']['keys']
        slide.notes_slide.notes_text_frame.text = store._encode(provenance)
    output = BytesIO(); deck.save(output)
    return output.getvalue()
